"""XI2 layers 1–2 on the XI1 fixtures with a fake context: the profile
layer re-validates and makes no model call, the rules layer maps the messy
layouts, names outside the context are dropped, and a second company's
names never appear. No database, no network, no model."""
import copy

import pytest

from app.schemas.mapping import MappingProposal
from app.services.mapping.engine import propose
from app.services.mapping.synonyms import part_of_header
from app.services.workbook import read_workbook
from ci.tests.conftest import FIXTURES
from ci.tests.mapping_fixtures import ACTIVITY_NAMES, OTHER_ACTIVITY_NAMES, ctx, lookup

XI1 = FIXTURES / "xi1"


def proposal_for(name, profiles=None, site_id=2):
    wb = read_workbook(XI1 / name)
    return propose(wb, 1, ctx(site_id), profiles or [], lookup())


def fields(sheet):
    return {f.key: f.column for f in sheet.fields}


# ---------------------------------------------------------------------------
# Rules
# ---------------------------------------------------------------------------


def test_single_sheet_monthly_maps_every_part_from_the_headers():
    s = proposal_for("01-single-sheet-monthly.xlsx").sheets[0]
    assert fields(s) == {"activity_name": "Activity", "activity_value": "Quantity", "activity_unit": "Unit", "supplier": None}
    assert s.period.kind == "column" and s.period.column == "Month"
    # "HSD" is a Stationary combustion name only, so the sheet's category follows its activities.
    assert s.category.kind == "fixed" and s.category.id == 3 and s.category.confidence == 0.85
    assert [(a.sheet_text, a.factor_name, a.source) for a in s.activities] == [("HSD", "HSD", "rules")]
    assert s.source == "rules" and s.confidence >= 0.8


def test_sheet_named_after_a_category_is_that_category():
    p = proposal_for("04-one-sheet-per-category.xlsx")
    assert [(s.sheet_name, s.category.id, s.category.confidence) for s in p.sheets] == [("Stationary combustion", 3, 1.0), ("Purchased electricity", 5, 1.0), ("Business travel", 6, 1.0)]
    travel = p.sheets[2]
    # The distance form's field claims "Distance (km)"; the generic value stays empty.
    assert fields(travel)["distance"] == "Distance (km)" and fields(travel)["activity_value"] is None
    assert fields(travel)["passengers"] == "Passengers"


def test_category_column_and_site_column_in_one_sheet():
    s = proposal_for("03-multi-site-one-sheet.xlsx").sheets[0]
    assert s.category.kind == "column" and s.category.column == "Category"
    assert [(v.text, v.id) for v in s.category.values] == [("Stationary combustion", 3)]
    assert s.site_column.column == "Plant"
    assert [(v.text, v.matches_site) for v in s.site_column.values] == [("Plant 2", True), ("Plant 9", False)]
    assert s.not_in_context == ["Plant 9"]
    assert s.site_id == 2


def test_yearly_sheets_are_read_as_cy_and_fy():
    cy = proposal_for("05-yearly-cy.xlsx").sheets[0]
    assert cy.period.model_dump() == {"kind": "yearly", "year_type": "CY", "year": 2024}
    assert cy.category.kind == "column" and [v.id for v in cy.category.values] == [3, 4, 5]
    fy = proposal_for("06-yearly-fy.xlsx").sheets[0]
    assert fy.period.model_dump() == {"kind": "yearly", "year_type": "FY", "year": 2025}  # FY 2024-25 is filed on Mar 31, 2025
    # "Scope" is not a category column: none of its values is a category name.
    assert fy.category.kind == "fixed"
    assert fy.confidence < 0.8


def test_months_as_columns_period_with_month_end_dates():
    s = proposal_for("07-months-as-columns.xlsx").sheets[0]
    assert s.layout == "months_as_columns" and s.period.kind == "months_as_columns"
    assert [c.date for c in s.period.columns][:3] == ["2025-01-31", "2025-02-28", "2025-03-31"]
    assert fields(s)["activity_value"] is None and fields(s)["activity_unit"] == "Unit"
    # Diesel belongs to two categories here, so the category stays for the user to confirm.
    assert s.confidence < 0.8


def test_electricity_bills_prefer_the_physical_quantity_over_the_amount():
    s = proposal_for("02-electricity-bills-date-column.xlsx").sheets[0]
    f = fields(s)
    assert f["activity_value"] == "Units consumed (kWh)" and f["supplier"] == "Supplier"
    assert s.period.column == "Bill date"


def test_mixed_units_value_column_is_claimed_with_low_confidence():
    s = proposal_for("08-mixed-units.xlsx").sheets[0]
    value = next(x for x in s.fields if x.key == "activity_value")
    assert value.column == "Consumption" and value.confidence < 0.8
    assert {a.sheet_text: a.factor_name for a in s.activities} == {"Diesel": "Diesel", "LPG": "LPG", "Coal": None}


def test_large_ledger_resolves_everything_above_threshold():
    s = proposal_for("12-large-ledger.csv").sheets[0]
    assert s.confidence >= 0.8 and s.category.kind == "column" and s.category.values[0].id == 4
    assert {a.sheet_text for a in s.activities} == {"CNG", "Diesel", "Petrol"}
    assert all(a.factor_name == a.sheet_text for a in s.activities)


def test_hostile_cells_change_nothing_but_the_text_they_hold():
    s = proposal_for("11-hostile-cell.xlsx").sheets[0]
    assert s.site_id == 2 and fields(s)["activity_value"] == "Qty"
    assert all(a.sheet_text == "Diesel" for a in s.activities)
    assert "Plant 9" not in s.not_in_context


def test_proposal_validates_against_the_contract_and_names_only_context_ids():
    for name in ("03-multi-site-one-sheet.xlsx", "05-yearly-cy.xlsx", "07-months-as-columns.xlsx"):
        p = proposal_for(name)
        MappingProposal.model_validate(p.model_dump())
        for s in p.sheets:
            assert s.site_id == 2
            ids = [s.category.id] if s.category.kind == "fixed" else [v.id for v in s.category.values if v.id]
            assert set(ids) <= {3, 4, 5, 6}
            for a in s.activities:
                assert a.category_id in (3, 4, 5, 6)
                assert a.factor_name is None or a.factor_name in ACTIVITY_NAMES[a.category_id]


def test_another_companys_names_never_appear():
    for name in ("01-single-sheet-monthly.xlsx", "03-multi-site-one-sheet.xlsx", "12-large-ledger.csv"):
        text = proposal_for(name).model_dump_json()
        for secret in ["Zed", "Process emissions", *OTHER_ACTIVITY_NAMES[30]]:
            assert secret not in text


def test_a_site_outside_the_context_is_refused():
    with pytest.raises(ValueError):
        proposal_for("01-single-sheet-monthly.xlsx", site_id=7)


# ---------------------------------------------------------------------------
# Profile
# ---------------------------------------------------------------------------


def saved_profile_for(name):
    """A profile confirmed on the fixture's own rules proposal, as XI3 would store it."""
    sheet = proposal_for(name).sheets[0]
    return {"id": "0f6a6e3c-1111-4aaa-8bbb-ccccccccdddd", "company_id": 1, "site_id": 2, "category_id": None, "fingerprint": sheet.fingerprint, "sheet_proposal": sheet.model_dump(), "uses": 1, "created_at": "2026-10-14T00:00:00Z", "updated_at": "2026-10-14T00:00:00Z"}


def test_profile_hit_maps_with_confidence_one_and_no_model_call():
    profile = saved_profile_for("07-months-as-columns.xlsx")
    profile["sheet_proposal"]["category"] = {"kind": "fixed", "id": 4, "label": "Mobile combustion", "confidence": 1.0, "source": "profile"}
    p = proposal_for("07-months-as-columns.xlsx", profiles=[profile])
    s = p.sheets[0]
    assert p.model.cached_profile_hits == 1 and p.model.calls == 0
    assert s.category.kind == "fixed" and s.category.id == 4 and s.category.source == "profile"
    assert s.period.kind == "months_as_columns" and s.confidence >= 0.8
    assert {f.key: f.source for f in s.fields}["activity_unit"] == "profile"


def test_profile_for_another_site_or_fingerprint_is_ignored():
    profile = saved_profile_for("01-single-sheet-monthly.xlsx")
    other_site = {**profile, "site_id": 9}
    other_fp = {**profile, "fingerprint": "0" * 64}
    p = proposal_for("01-single-sheet-monthly.xlsx", profiles=[other_site, other_fp])
    assert p.model.cached_profile_hits == 0 and p.sheets[0].source == "rules"


def test_profile_parts_that_vanished_from_the_context_fall_to_the_rules():
    profile = saved_profile_for("01-single-sheet-monthly.xlsx")
    profile["sheet_proposal"]["category"] = {"kind": "fixed", "id": 99, "label": "Gone", "confidence": 1.0, "source": "profile"}
    profile["sheet_proposal"]["activities"] = [{"sheet_text": "HSD", "category_id": 3, "factor_name": "Vanished factor", "confidence": 1.0, "source": "profile", "save_mapping": False}]
    s = proposal_for("01-single-sheet-monthly.xlsx", profiles=[profile]).sheets[0]
    assert s.category.id == 3 and s.category.source == "rules"
    assert s.activities[0].factor_name == "HSD" and s.activities[0].source == "rules"
    assert s.source == "mixed"


def test_profile_fixed_period_is_not_reused_for_a_new_file():
    profile = saved_profile_for("05-yearly-cy.xlsx")
    profile["sheet_proposal"]["period"] = {"kind": "yearly", "year_type": "CY", "year": 2019}
    s = proposal_for("05-yearly-cy.xlsx", profiles=[profile]).sheets[0]
    assert s.period.model_dump() == {"kind": "yearly", "year_type": "CY", "year": 2024}


# ---------------------------------------------------------------------------
# Synonyms
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "header, part",
    [("Qty", "activity_value"), ("Units consumed (kWh)", "activity_value"), ("UOM", "activity_unit"), ("Bill date", "period"), ("Fuel type", "activity_name"), ("Plant", "site"), ("Emission category", "activity_name")],
)
def test_part_of_header(header, part):
    assert part in part_of_header(header)


def test_unknown_header_means_nothing():
    assert part_of_header("Remarks") == [] and part_of_header("") == []


# ---------------------------------------------------------------------------
# The user's category choice (confirm screen re-asks for one sheet)
# ---------------------------------------------------------------------------

def test_a_chosen_category_is_taken_as_resolved_and_its_own_fields_are_mapped():
    wb = read_workbook(XI1 / "01-single-sheet-monthly.xlsx")
    name = wb.readable_sheets[0].name
    s = propose(wb, 1, ctx(2), [], lookup(), sheet_categories={name: 6}).sheets[0]
    # Business travel (6) instead of the rules' Stationary combustion (3); its form wants a distance
    assert s.category.kind == "fixed" and s.category.id == 6 and s.category.confidence == 1.0
    assert "distance" in {f.key for f in s.fields}


def test_only_the_named_sheets_are_proposed():
    wb = read_workbook(XI1 / "04-one-sheet-per-category.xlsx")
    names = [sh.name for sh in wb.readable_sheets]
    assert len(names) > 1
    p = propose(wb, 1, ctx(2), [], lookup(), only_sheets=[names[1]])
    assert [sh.sheet_name for sh in p.sheets] == [names[1]]


def test_a_chosen_category_outside_the_context_is_refused():
    wb = read_workbook(XI1 / "01-single-sheet-monthly.xlsx")
    with pytest.raises(ValueError):
        propose(wb, 1, ctx(2), [], lookup(), sheet_categories={wb.readable_sheets[0].name: 30})


def test_a_saved_profile_of_another_category_does_not_override_the_choice():
    wb = read_workbook(XI1 / "01-single-sheet-monthly.xlsx")
    first = propose(wb, 1, ctx(2), [], lookup()).sheets[0]
    profile = {"fingerprint": first.fingerprint, "site_id": 2, "category_id": 3, "sheet_proposal": first.model_dump()}
    s = propose(wb, 1, ctx(2), [profile], lookup(), sheet_categories={first.sheet_name: 5}).sheets[0]
    assert s.category.id == 5 and s.source != "profile"


# ---------------------------------------------------------------------------
# Unit from the header; a "Unit" column that holds no units
# ---------------------------------------------------------------------------

def _ac_topups(tmp_path):
    from openpyxl import Workbook
    import datetime as dt
    wb = Workbook(); ws = wb.active; ws.title = "AC top-ups"
    ws.append(["Service date", "Unit", "Gas", "Charged (kg)", "Vendor"])
    ws.append([dt.date(2025, 5, 14), "Admin block AHU", "R410A", 6.5, "CoolServ"])
    ws.append([dt.date(2025, 9, 2), "Cold room 1", "R410A", 4.0, "CoolServ"])
    ws.append([dt.date(2025, 11, 20), "QC lab split AC", "R410A", 1.2, "Local tech"])
    p = tmp_path / "ac.xlsx"; wb.save(p)
    return read_workbook(p)


def test_a_unit_column_holding_equipment_names_is_not_the_unit_and_the_header_unit_needs_no_column(tmp_path):
    from app.services.mapping.prompt import needed_parts
    from app.services.mapping.layers import Draft, rules_layer
    wb = _ac_topups(tmp_path)
    s = propose(wb, 1, ctx(2), [], lookup(), sheet_categories={"AC top-ups": 3}).sheets[0]
    unit = next(f for f in s.fields if f.key == "activity_unit")
    assert unit.column is None and unit.skipped is True and unit.confidence == 1.0
    assert fields(s)["activity_value"] == "Charged (kg)"
    # the model is not asked for the unit
    d = Draft(sheet=wb.readable_sheets[0], ctx=ctx(2))
    rules_layer(d, lookup())
    assert d.header_unit == "kg" and "activity_unit" not in needed_parts(d).get("fields", [])
    # and every row gets the header unit
    rows = rows_from_proposal_for(wb, s)
    assert {r["activity_unit"] for r in rows} == {"kg"}


def rows_from_proposal_for(wb, sheet):
    from app.services.mapping.rows import rows_from_proposal
    p = MappingProposal.model_validate({"document_id": 1, "site_id": 2, "sheets": [sheet.model_dump()], "model": {"name": "none", "calls": 0, "cached_profile_hits": 0}})
    return rows_from_proposal(wb, p).rows


def _quantity_named_ctx(calc_mode="single"):
    """Category 3 as Glochem's form has it: the quantity column is a required number field named "Fuel consumption"."""
    from app.services.mapping.context import context_view
    from ci.tests.mapping_fixtures import CONTEXT
    raw = copy.deepcopy(CONTEXT)
    cat = next(c for s in raw["sites"] if s["site_id"] == 2 for c in s["categories"] if c["category_id"] == 3)
    cat["form"] = {"calc_mode": calc_mode, "fields": [{"key": "Fuel consumption", "label": "Fuel consumption", "type": "number", "required": True}, {"key": "supplier", "label": "Supplier", "type": "text", "required": False}]}
    return context_view(raw, 2)


def _diesel_issues(tmp_path, with_own_column=False):
    from openpyxl import Workbook
    import datetime as dt
    wb = Workbook(); ws = wb.active; ws.title = "Diesel issues"
    ws.append(["Date", "Fuel", "Qty", "UoM"] + (["Fuel consumption"] if with_own_column else []))
    for day, qty in ((3, 1200), (17, 600), (28, 450)):
        ws.append([dt.date(2025, 7, day), "HSD", qty, "L"] + ([qty * 2] if with_own_column else []))
    p = tmp_path / "diesel.xlsx"; wb.save(p)
    return read_workbook(p)


def test_a_single_mode_forms_required_number_field_takes_the_quantity_column(tmp_path):
    from app.services.mapping.prompt import needed_parts
    from app.services.mapping.layers import Draft, rules_layer
    wb = _diesel_issues(tmp_path)
    s = propose(wb, 1, _quantity_named_ctx(), [], lookup(), sheet_categories={"Diesel issues": 3}).sheets[0]
    fc = next(f for f in s.fields if f.key == "Fuel consumption")
    assert fields(s)["activity_value"] == "Qty" and fc.column == "Qty" and fc.skipped is False and fc.confidence >= 0.8
    # nothing left for the model to find
    d = Draft(sheet=wb.readable_sheets[0], ctx=_quantity_named_ctx())
    rules_layer(d, lookup())
    assert "Fuel consumption" not in needed_parts(d).get("fields", [])
    # the rows carry the quantity under the form's name too, as a typed entry does
    rows = rows_from_proposal_for(wb, s)
    assert [(r["activity_value"], r["fields"]["Fuel consumption"]) for r in rows] == [(1200, "1200"), (600, "600"), (450, "450")]


def test_the_quantity_rule_leaves_an_own_column_and_spec_forms_alone(tmp_path):
    # a sheet with its own "Fuel consumption" column keeps it
    s = propose(_diesel_issues(tmp_path, with_own_column=True), 1, _quantity_named_ctx(), [], lookup(), sheet_categories={"Diesel issues": 3}).sheets[0]
    assert fields(s)["Fuel consumption"] == "Fuel consumption"
    # a form with a calculation spec multiplies its own fields: no copy of the quantity
    s = propose(_diesel_issues(tmp_path), 1, _quantity_named_ctx("per_method"), [], lookup(), sheet_categories={"Diesel issues": 3}).sheets[0]
    assert fields(s)["Fuel consumption"] is None


def _consumption_ctx():
    """Category 3 with Midal's form: one required number field "Consumption", no calculation spec."""
    from app.services.mapping.context import context_view
    from ci.tests.mapping_fixtures import CONTEXT
    raw = copy.deepcopy(CONTEXT)
    cat = next(c for s in raw["sites"] if s["site_id"] == 2 for c in s["categories"] if c["category_id"] == 3)
    cat["form"] = {"calc_mode": None, "fields": [{"key": "Fuel Type", "label": "Fuel Type", "type": "select", "required": False}, {"key": "Consumption", "label": "Consumption", "type": "number", "required": True}]}
    return context_view(raw, 2)


def test_the_quantity_takes_the_forms_named_consumption_column_not_a_serial_number(tmp_path):
    """Midal's Scope 1 template: S.No and Total Consumption are both numbers. The form's
    "Consumption" field finds "Total Consumption" by name; the quantity is that column too."""
    from openpyxl import Workbook
    import datetime as dt
    wb = Workbook(); ws = wb.active; ws.title = "Scope-1 Stationary Combustion"
    ws.append(["S.No", "Month/Period", "Source of Combustion", "Consumption Unit(Kg/MT/L)", "Total Consumption"])
    for i, (m, qty) in enumerate(((1, 1200), (2, 600), (3, 450)), start=1):
        ws.append([i, dt.date(2025, m, 1), "HSD", "L", qty])
    p = tmp_path / "scope1.xlsx"; wb.save(p)
    book = read_workbook(p)
    s = propose(book, 1, _consumption_ctx(), [], lookup(), sheet_categories={"Scope-1 Stationary Combustion": 3}).sheets[0]
    value = next(f for f in s.fields if f.key == "activity_value")
    assert fields(s)["Consumption"] == "Total Consumption"
    assert value.column == "Total Consumption" and value.confidence >= 0.8
    rows = rows_from_proposal_for(book, s)
    assert [r["activity_value"] for r in rows] == [1200, 600, 450]


def test_a_low_confidence_quantity_guess_gives_way_to_the_named_field(tmp_path):
    from app.services.mapping.layers import Draft, Part, quantity_fields, rules_layer
    from app.services.mapping.context import VALUE_KEY
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Fuel"
    ws.append(["S.No", "Month", "Unit", "Total Consumption"])
    for i, m in enumerate(("Jan-25", "Feb-25", "Mar-25"), start=1):
        ws.append([i, m, "L", 100 * i])
    p = tmp_path / "f.xlsx"; wb.save(p)
    sheet = read_workbook(p).readable_sheets[0]
    d = Draft(sheet=sheet, ctx=_consumption_ctx())
    rules_layer(d, lookup())
    # what the model layer did on Shyam's file: a guess at S.No under 0.8
    d.fields[VALUE_KEY] = Part("S.No", 0.6, "llm")
    quantity_fields(d)
    assert d.fields[VALUE_KEY].value == "Total Consumption"
    # a confident quantity is never replaced
    d.fields[VALUE_KEY] = Part("S.No", 0.9, "llm")
    quantity_fields(d)
    assert d.fields[VALUE_KEY].value == "S.No"


# ---------------------------------------------------------------------------
# Review findings (2026-10-14)
# ---------------------------------------------------------------------------


def test_a_profile_hit_applies_the_saved_activity_matches():
    profile = saved_profile_for("07-months-as-columns.xlsx")
    sp = profile["sheet_proposal"]
    sp["category"] = {"kind": "fixed", "id": 4, "label": "Mobile combustion", "confidence": 1.0, "source": "profile"}
    sp["activities"] = [{"sheet_text": "Grid electricity", "category_id": 4, "factor_name": "CNG", "confidence": 1.0, "source": "profile", "save_mapping": False}]
    s = proposal_for("07-months-as-columns.xlsx", profiles=[profile]).sheets[0]
    match = {a.sheet_text: a for a in s.activities}["Grid electricity"]
    assert (match.factor_name, match.source, match.confidence) == ("CNG", "profile", 1.0)


@pytest.mark.parametrize(
    "text, ok",
    [("Plant 2", True), ("plant-2", True), ("Plant 2 site", True), ("Plant 22", False), ("Plant 2 - Hyd", False), ("Plant 9", False), ("2", True)],
)
def test_a_site_cell_names_the_site_only_when_it_says_the_same(text, ok):
    from app.services.mapping.match import names_site

    assert names_site(text, "Plant 2", ["Plant 9"]) is ok
    assert names_site("Pune 2", "Pune", []) is False and names_site("Pune plant", "Pune", []) is True


def test_a_cell_longer_than_the_contract_allows_does_not_break_the_proposal():
    import io

    from openpyxl import Workbook as XlsxWorkbook

    from app.services.workbook import read_workbook_bytes

    wb = XlsxWorkbook()
    ws = wb.active
    ws.title = "Fuel"
    ws.append(["Month", "Fuel", "Qty", "Unit"])
    ws.append(["2025-01", "Diesel", 10, "litre"])
    ws.append(["2025-02", "x" * 400, 12, "litre"])
    buf = io.BytesIO()
    wb.save(buf)
    p = propose(read_workbook_bytes(buf.getvalue(), "xlsx"), 1, ctx(), [], lookup())
    MappingProposal.model_validate(p.model_dump())
    assert [a.sheet_text for a in p.sheets[0].activities] == ["Diesel"]


def test_on_a_months_as_columns_sheet_the_quantity_field_is_the_month_cells_not_the_total(tmp_path):
    from openpyxl import Workbook
    from app.services.mapping.layers import Draft, rules_layer
    from app.services.mapping.prompt import needed_parts
    wb = Workbook(); ws = wb.active; ws.title = "Diesel by month"
    ws.append(["Diesel issued (litre) FY 2025-26"]); ws.append([])
    ws.append(["Fuel", "Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "FY Total"])
    ws.append(["HSD", 100, 110, 120, 130, 140, 150, 750])
    ws.append(["LDO", 10, 11, 12, 13, 14, 15, 75])
    p = tmp_path / "pivot.xlsx"; wb.save(p)
    book = read_workbook(p)
    s = propose(book, 1, _quantity_named_ctx(), [], lookup(), sheet_categories={"Diesel by month": 3}).sheets[0]
    fc = next(f for f in s.fields if f.key == "Fuel consumption")
    assert s.layout == "months_as_columns" and fc.column is None and fc.skipped is True
    d = Draft(sheet=book.readable_sheets[0], ctx=_quantity_named_ctx())
    rules_layer(d, lookup())
    assert "Fuel consumption" not in needed_parts(d).get("fields", [])
    rows = rows_from_proposal_for(book, s)
    assert sorted(r["activity_value"] for r in rows) == sorted([100, 110, 120, 130, 140, 150, 10, 11, 12, 13, 14, 15])
    assert not any("Fuel consumption" in (r.get("fields") or {}) for r in rows)


@pytest.mark.parametrize(
    "text,year",
    [("FY 2025-26", 2026), ("FY 2024-25", 2025), ("FY2024/2025", 2025), ("FY 2029-30", 2030), ("FY25", 2025), ("FY 2025", 2025)],
)
def test_an_fy_label_gives_the_year_the_fy_ends(text, year):
    from app.services.mapping.layers import _FY_RE, _fy_end_year

    assert _fy_end_year(_FY_RE.search(text)) == year
