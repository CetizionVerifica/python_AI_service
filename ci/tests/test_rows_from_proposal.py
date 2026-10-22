"""XI2 ``rows_from_proposal``: a confirmed proposal plus the workbook give
one ImportRow per kept sheet row (or per filled month cell), with the row
ref, period, wording and unit as written, provenance in ``meta``, and local
issues for rows this service already knows ESG-lite would refuse."""
import pytest

from app.schemas.mapping import MappingProposal
from app.services.mapping.engine import propose
from app.services.mapping.rows import rows_from_proposal
from app.services.workbook import read_workbook
from ci.tests.conftest import FIXTURES
from ci.tests.mapping_fixtures import ctx, lookup

XI1 = FIXTURES / "xi1"


def planned(name, edit=None):
    wb = read_workbook(XI1 / name)
    p = propose(wb, 1, ctx(), [], lookup())
    if edit:
        p = MappingProposal.model_validate(edit(p.model_dump()))
    return wb, p, rows_from_proposal(wb, p)


def test_one_row_per_kept_sheet_row_with_ref_period_and_text_as_written():
    _, _, sr = planned("01-single-sheet-monthly.xlsx")
    assert len(sr.rows) == 8 and sr.local_issues == []
    first = sr.rows[0]
    assert first["row_ref"] == "Diesel!A2" and first["site_id"] == 2 and first["category_id"] == 3
    assert first["period"] == {"kind": "monthly", "date": "2025-01-31"}
    assert first["activity_name"] == "HSD" and first["activity_unit"] == "Litre" and first["activity_value"] == 1200
    assert first["meta"] == {"source": "rules", "confidence": 0.85}
    assert "total" not in first and "factor" not in first


def test_rows_for_another_site_are_kept_back_as_site_not_allowed():
    _, _, sr = planned("03-multi-site-one-sheet.xlsx")
    assert len(sr.rows) == 8
    assert [(i.code, i.row_ref) for i in sr.local_issues] == [("SITE_NOT_ALLOWED", "All plants!A3"), ("SITE_NOT_ALLOWED", "All plants!A6"), ("SITE_NOT_ALLOWED", "All plants!A9"), ("SITE_NOT_ALLOWED", "All plants!A12")]
    assert all(r["site_id"] == 2 for r in sr.rows)
    assert sr.local_issues[0].as_planned_row()["state"] == "issue"


def test_months_as_columns_make_one_row_per_filled_cell():
    _, p, sr = planned("07-months-as-columns.xlsx")
    assert len(sr.rows) == 35  # 3 activities × 12 months, one blank cell
    assert sr.rows[0]["row_ref"] == "Pivot!C2" and sr.rows[0]["period"] == {"kind": "monthly", "date": "2025-01-31"}
    assert sr.rows[11]["row_ref"] == "Pivot!N2" and sr.rows[11]["period"]["date"] == "2025-12-31"
    blank = sr.local_issues[0]
    assert (blank.code, blank.row_ref, blank.severity) == ("PIVOT_EMPTY_CELL", "Pivot!G4", "info")
    assert blank.as_planned_row()["state"] == "skipped"


def test_yearly_period_is_carried_on_every_row():
    _, _, sr = planned("06-yearly-fy.xlsx")
    assert [r["period"] for r in sr.rows] == [{"kind": "yearly", "year_type": "FY", "year": 2025}] * 3  # FY 2024-25 ends Mar 31, 2025
    assert [r["activity_unit"] for r in sr.rows] == ["L", "kg", "kWh"]


def test_units_inside_value_cells_and_headers_are_used():
    _, _, sr = planned("08-mixed-units.xlsx")
    assert [(r["activity_value"], r["activity_unit"]) for r in sr.rows] == [(1200.0, "litre"), (1.5, "KL"), (400.0, "kg"), (0.4, "MT"), (12.0, "tonnes")]
    assert sr.rows[0]["period"]["date"] == "2025-04-01"  # day-first
    _, _, sr = planned("04-one-sheet-per-category.xlsx")
    electricity = [r for r in sr.rows if r["category_id"] == 5]
    assert electricity and all(r["activity_unit"] == "kwh" for r in electricity)  # from the "kWh" header
    travel = [r for r in sr.rows if r["category_id"] == 6]
    assert travel and travel[0]["fields"] == {"distance": "1700", "passengers": "1"} and travel[0]["activity_unit"] == "km"
    assert "activity_value" not in travel[0]


def test_category_column_rows_without_a_category_are_kept_back():
    def edit(d):
        d["sheets"][0]["category"]["values"][0]["id"] = None
        return d

    _, _, sr = planned("05-yearly-cy.xlsx", edit)
    assert [i.code for i in sr.local_issues] == ["CATEGORY_NOT_ASSIGNED"] and len(sr.rows) == 2


def test_excluded_sheets_and_skipped_fields_are_left_out():
    def edit(d):
        d["sheets"][1]["excluded"] = True
        for f in d["sheets"][0]["fields"]:
            if f["key"] == "supplier":
                f["column"] = None
                f["skipped"] = True
        return d

    _, _, sr = planned("04-one-sheet-per-category.xlsx", edit)
    assert {r["category_id"] for r in sr.rows} == {3, 6}
    assert all("fields" not in r for r in sr.rows if r["category_id"] == 3)


def test_a_changed_file_is_refused_by_fingerprint():
    wb, p, _ = planned("01-single-sheet-monthly.xlsx")
    other = read_workbook(XI1 / "10-subtotal-rows.xlsx")
    other.sheets[0].name = "Diesel"
    with pytest.raises(ValueError):
        rows_from_proposal(other, p)


def test_large_ledger_maps_every_row_within_the_plan_limit():
    _, _, sr = planned("12-large-ledger.csv")
    assert len(sr.rows) == 15000 and sr.total_rows == 15000
    assert sr.rows[-1]["row_ref"] == "12-large-ledger!A15001"


def _pivot_with_title(tmp_path, title):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Power"
    ws.append([title]); ws.append([])
    ws.append(["Meter", "Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "FY Total"])
    ws.append(["Main incomer", 100, 110, 120, 130, 140, 150, 750])
    ws.append(["Utility block", 10, 11, 12, 13, 14, 15, 75])
    p = tmp_path / "power.xlsx"; wb.save(p)
    return read_workbook(p)


def test_a_unit_named_only_in_the_title_is_the_unit_of_every_month_cell(tmp_path):
    wb = _pivot_with_title(tmp_path, "Electricity consumption (kWh) - FY 2025-26")
    assert wb.readable_sheets[0].title_units == ["kwh"]
    proposal = propose(wb, 1, ctx(2), [], lookup(), sheet_categories={"Power": 5})
    rows = rows_from_proposal(wb, proposal)
    assert len(rows.rows) == 12  # 2 meters x 6 months; the FY Total column is not a month
    assert {r["activity_unit"] for r in rows.rows} == {"kwh"}
    assert not [i for i in rows.local_issues if i.code == "MISSING_UNIT"]


def test_two_units_in_the_title_are_ambiguous_and_give_no_unit(tmp_path):
    wb = _pivot_with_title(tmp_path, "Energy (kWh) and diesel (litre)")
    proposal = propose(wb, 1, ctx(2), [], lookup(), sheet_categories={"Power": 5})
    rows = rows_from_proposal(wb, proposal)
    assert {r.get("activity_unit") for r in rows.rows} <= {"", None} or not rows.rows


# ---------------------------------------------------------------------------
# Review findings (2026-10-14)
# ---------------------------------------------------------------------------


def _pivot_workbook(cells):
    import io

    from openpyxl import Workbook as XlsxWorkbook

    from app.services.workbook import read_workbook_bytes

    wb = XlsxWorkbook()
    ws = wb.active
    ws.title = "Pivot"
    months = ["Jan-25", "Feb-25", "Mar-25", "Apr-25", "May-25", "Jun-25"]
    ws.append(["Activity", "Unit"] + months)
    for row in cells:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return read_workbook_bytes(buf.getvalue(), "xlsx")


def _pivot_rows(cells, unit_column=True):
    wb = _pivot_workbook(cells)
    p = propose(wb, 1, ctx(), [], lookup())
    if not unit_column:
        d = p.model_dump()
        for f in d["sheets"][0]["fields"]:
            if f["key"] == "activity_unit":
                f["column"], f["skipped"] = None, True
        p = MappingProposal.model_validate(d)
    return rows_from_proposal(wb, p)


def test_a_unit_seen_in_one_month_cell_is_not_the_unit_of_the_others():
    sr = _pivot_rows([["Diesel", None, "5 t", 6, 7, 8, 9, 10], ["Petrol", None, 1, 2, 3, 4, 5, 6]], unit_column=False)
    assert [(r["row_ref"], r["activity_unit"]) for r in sr.rows if r["row_ref"].startswith("Pivot!C")] == [("Pivot!C2", "t")]
    others = [i for i in sr.local_issues if i.code == "MISSING_UNIT"]
    assert len(others) == 11


def test_months_as_columns_cells_that_are_not_numbers_are_flagged_and_dashes_are_blank():
    sr = _pivot_rows([["Diesel", "litre", 10, "NA", "-", "abc", "35 m3", 12], ["Petrol", "litre", 1, 2, 3, 4, 5, 6]])
    by_ref = {r["row_ref"]: r for r in sr.rows}
    issues = {i.row_ref: (i.code, i.severity) for i in sr.local_issues}
    assert issues["Pivot!D2"] == ("PIVOT_EMPTY_CELL", "info") and issues["Pivot!E2"] == ("PIVOT_EMPTY_CELL", "info")
    assert issues["Pivot!F2"] == ("VALUE_NOT_NUMERIC", "error")
    assert by_ref["Pivot!G2"]["activity_value"] == 35.0
    assert all("fields" not in r or "activity_value" not in r["fields"] for r in sr.rows)


def test_a_row_with_a_cell_too_long_for_esg_lite_is_kept_back():
    sr = _pivot_rows([["Diesel", "litre", 1, 2, 3, 4, 5, 6], ["y" * 301, "litre", 1, 2, 3, 4, 5, 6]])
    assert {i.code for i in sr.local_issues} == {"TEXT_TOO_LONG"} and len(sr.rows) == 6


# ---------------------------------------------------------------------------
# #183 P1: when a row's unit readings disagree, the one its activity accepts
# ---------------------------------------------------------------------------

from app.services.mapping.context import AcceptedUnitsLookup  # noqa: E402
from app.services.mapping.rows import pick_unit  # noqa: E402


def test_pick_unit_keeps_the_first_reading_unless_only_a_later_one_is_accepted():
    asked = []

    def accepted():
        asked.append(1)
        return ["kg", "tonne"]

    assert pick_unit([None, "", "Litre", "kg"]) == "Litre"  # no list: the first reading, as before
    assert pick_unit(["litre", "kg"], accepted) == "kg"  # the unit column disagrees; the header's kg is accepted
    assert pick_unit(["kg", "tonne"], accepted) == "kg"  # both accepted: the first
    assert pick_unit(["MT", "kg"], accepted) == "MT"  # MT is the engine's tonne, so it counts as accepted
    assert pick_unit(["litre", "kWh"], accepted) == "litre"  # none accepted: the first, the server says why
    asked.clear()
    assert pick_unit(["kg", "KG", " kg "], accepted) == "kg" and asked == []  # readings agree: nothing is asked
    assert pick_unit([None, None]) == ""


def _sheet_with_unit_column(tmp_path, units):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Fuel"
    ws.append(["Date", "Fuel", "Quantity (kg)", "Unit"])
    for i, u in enumerate(units):
        ws.append([f"2025-0{i + 1}-28", "LPG", 100 + i, u])
    p = tmp_path / "fuel.xlsx"; wb.save(p)
    return read_workbook(p)


def test_a_unit_column_that_disagrees_with_the_quantity_header_takes_the_accepted_reading(tmp_path):
    wb = _sheet_with_unit_column(tmp_path, ["litre", "kg"])
    proposal = propose(wb, 1, ctx(2), [], lookup(), sheet_categories={"Fuel": 3})
    plain = rows_from_proposal(wb, proposal)
    assert [r["activity_unit"] for r in plain.rows] == ["litre", "kg"]  # without the list: the unit column, as before

    calls = []

    def accepted(category_id, name):
        calls.append((category_id, name))
        return ["kg", "tonne"] if (category_id, name) == (3, "LPG") else None

    hinted = rows_from_proposal(wb, proposal, accepted)
    assert [r["activity_unit"] for r in hinted.rows] == ["kg", "kg"]
    assert calls == [(3, "LPG")]  # asked once, for the row whose readings disagree


def test_accepted_units_are_read_by_name_through_the_activity_search_and_cached():
    class Search:
        def __init__(self):
            self.calls = 0

        def options(self, category_id, q=""):
            self.calls += 1
            return [{"name": "LPG", "accepted_units": ["kg", "tonne"]}, {"name": "LPG cylinders", "accepted_units": ["cylinder"]}]

    search = Search()
    accepted = AcceptedUnitsLookup(search)
    assert accepted(3, "lpg") == ["kg", "tonne"]
    assert accepted(3, "LPG") == ["kg", "tonne"] and search.calls == 1
    assert accepted(3, "Coal") is None  # not one of the category's names
