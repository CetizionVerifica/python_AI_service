"""XI1: the workbook reader gives every fixture layout the expected
description (golden JSON per fixture, ``UPDATE_GOLDEN=1`` refreshes), the
fingerprint ignores sheet names, row counts and column order, and nothing in
a cell is interpreted. No database, no network, no model."""
import io
import json
import os
from datetime import datetime

import pytest
from openpyxl import Workbook as XlsxWorkbook

from app.services.workbook import fingerprint_of, month_of_header, normalise_header, read_workbook, read_workbook_bytes
from ci.tests.conftest import FIXTURES

XI1 = FIXTURES / "xi1"
FIXTURE_FILES = sorted(p.name for p in XI1.iterdir() if p.suffix in {".xlsx", ".csv"})


def xlsx_bytes(build) -> bytes:
    wb = XlsxWorkbook()
    build(wb)
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Golden descriptions
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("name", FIXTURE_FILES)
def test_fixture_matches_golden_description(name):
    wb = read_workbook(XI1 / name)
    actual = wb.describe()
    golden_path = XI1 / "golden" / (name + ".json")
    if os.environ.get("UPDATE_GOLDEN") == "1":
        golden_path.parent.mkdir(exist_ok=True)
        golden_path.write_text(json.dumps(actual, indent=2, ensure_ascii=False) + "\n")
    expected = json.loads(golden_path.read_text())
    assert actual == expected, f"{name}: description changed (UPDATE_GOLDEN=1 to accept)"


def test_twelve_layouts_are_covered():
    assert len(FIXTURE_FILES) == 12


# ---------------------------------------------------------------------------
# Layouts
# ---------------------------------------------------------------------------


def sheet(name, index=0):
    return read_workbook(XI1 / name).sheets[index]


def test_banner_rows_above_the_header_are_skipped():
    s = sheet("02-electricity-bills-date-column.xlsx")
    assert s.header_row == 3 and s.header_depth == 1
    assert [c.name for c in s.columns] == ["Bill date", "Meter no", "Supplier", "Units consumed (kWh)", "Amount (INR)"]
    assert s.row_refs[0] == "Bills!A4"
    assert s.columns[0].dtype == "date" and s.columns[3].unit_tokens == ["kwh"]


def test_every_sheet_is_read_not_only_the_first():
    wb = read_workbook(XI1 / "04-one-sheet-per-category.xlsx")
    assert [s.name for s in wb.sheets] == ["Stationary combustion", "Purchased electricity", "Business travel"]
    assert all(s.skipped_reason is None for s in wb.sheets)


def test_total_subtotal_and_blank_rows_are_dropped_with_their_numbers():
    s = sheet("10-subtotal-rows.xlsx")
    assert s.data_rows == 6
    assert [(d.row, d.reason) for d in s.dropped_rows] == [(5, "total"), (9, "total"), (10, "blank"), (11, "total")]
    assert all(not r["cells"]["Month"].lower().startswith(("sub", "grand")) for r in s.rows)


def test_grand_total_row_in_yearly_sheet_is_dropped():
    assert [(d.row, d.reason) for d in sheet("06-yearly-fy.xlsx").dropped_rows] == [(5, "total")]


def test_months_as_columns_layout_with_month_end_dates():
    s = sheet("07-months-as-columns.xlsx")
    assert s.layout == "months_as_columns"
    months = [c for c in s.columns if c.month]
    assert len(months) == 12
    assert months[0].month == {"month": 1, "year": 2025, "date": "2025-01-31"}
    assert months[1].month["date"] == "2025-02-28"
    # The blank May cell is kept blank (XI2 raises PIVOT_EMPTY_CELL), the total row is dropped.
    assert s.rows[2]["cells"]["May-25"] == ""
    assert s.data_rows == 3


def test_two_row_merged_header_is_joined():
    s = sheet("09-two-row-merged-header.xlsx")
    assert s.header_row == 1 and s.header_depth == 2
    assert [c.name for c in s.columns] == ["Site", "Period", "Electricity / kWh", "Electricity / INR", "Diesel / Litre", "Diesel / INR"]
    assert s.row_refs[0] == "Energy!A3"
    assert s.columns[4].unit_tokens == ["litre"]


def test_mixed_units_inside_cells_are_listed_as_unit_tokens():
    s = sheet("08-mixed-units.xlsx")
    col = next(c for c in s.columns if c.name == "Consumption")
    assert col.dtype == "text"
    assert set(col.unit_tokens) == {"litre", "kl", "kg", "mt", "tonnes"}
    assert s.rows[0]["cells"]["Consumption"] == "1,200 litre"  # as written


def test_multi_site_sheet_keeps_the_site_column_values():
    s = sheet("03-multi-site-one-sheet.xlsx")
    plant = next(c for c in s.columns if c.name == "Plant")
    assert plant.distinct == 2 and set(plant.samples) == {"Plant 2", "Plant 9"}


def test_hostile_cells_are_carried_as_text_and_never_interpreted():
    s = sheet("11-hostile-cell.xlsx")
    notes = [r["cells"]["Note"] for r in s.rows]
    assert notes[0].startswith("IGNORE ALL PREVIOUS INSTRUCTIONS")
    assert notes[2] == "<script>alert(1)</script>"
    # A formula has no cached value in a fresh file: it reads as blank, never as a link.
    assert notes[1] == ""


def test_large_csv_ledger_reads_every_row():
    s = sheet("12-large-ledger.csv")
    assert s.data_rows == 15000 and s.row_refs[-1] == "12-large-ledger!A15001"
    assert s.columns[0].dtype == "date" and s.columns[4].dtype == "number"
    assert len(s.describe()["row_refs_head"]) == 3 and "rows" not in s.describe()


def test_yearly_sheets_keep_their_header_and_rows():
    cy = sheet("05-yearly-cy.xlsx")
    assert cy.header_row == 2 and cy.data_rows == 3
    fy = sheet("06-yearly-fy.xlsx")
    assert [c.name for c in fy.columns][1] == "FY 2024-25 (Apr-Mar)"


# ---------------------------------------------------------------------------
# Fingerprint
# ---------------------------------------------------------------------------


def test_fingerprint_is_stable_across_sheet_rename_row_count_and_column_order():
    def one(wb):
        ws = wb.active
        ws.title = "March export"
        ws.append(["Month", "Activity", "Qty (L)", "Unit"])
        for i in range(3):
            ws.append([f"2025-0{i + 1}", "Diesel", 10 + i, "litre"])

    def two(wb):
        ws = wb.active
        ws.title = "April export"
        ws.append(["Unit", "QTY  (l)", "activity", "MONTH"])
        for i in range(40):
            ws.append(["litre", 10 + i, "Diesel", f"2025-0{i % 9 + 1}"])

    a = read_workbook_bytes(xlsx_bytes(one), "xlsx").sheets[0]
    b = read_workbook_bytes(xlsx_bytes(two), "xlsx").sheets[0]
    assert a.fingerprint == b.fingerprint
    assert a.fingerprint == fingerprint_of(["Month", "Activity", "Qty (L)", "Unit"], 1)


def test_fingerprint_changes_with_a_header_or_the_depth():
    assert fingerprint_of(["Month", "Qty"], 1) != fingerprint_of(["Month", "Quantity"], 1)
    assert fingerprint_of(["Month", "Qty"], 1) != fingerprint_of(["Month", "Qty"], 2)


def test_normalise_header_drops_case_punctuation_and_spacing():
    assert normalise_header("  Qty.  (L) ") == "qty l"
    assert normalise_header("Units consumed (kWh)") == "units consumed kwh"


# ---------------------------------------------------------------------------
# Edge cases
# ---------------------------------------------------------------------------


def test_hidden_sheets_and_tiny_sheets_are_listed_with_a_reason():
    def build(wb):
        ws = wb.active
        ws.title = "Data"
        ws.append(["Month", "Qty"])
        ws.append(["Jan-25", 1])
        ws.append(["Feb-25", 2])
        hidden = wb.create_sheet("Lookup")
        hidden.sheet_state = "hidden"
        hidden.append(["a", "b"])
        hidden.append([1, 2])
        hidden.append([3, 4])
        tiny = wb.create_sheet("Notes")
        tiny.append(["Prepared by", "Finance"])
        tiny.append(["Checked", 1])

    wb = read_workbook_bytes(xlsx_bytes(build), "xlsx")
    assert [s.skipped_reason for s in wb.sheets] == [None, "hidden", "fewer than 2 data rows"]
    assert [s.name for s in wb.readable_sheets] == ["Data"]


def test_blank_and_duplicate_headers_get_usable_names():
    def build(wb):
        ws = wb.active
        ws.append(["Month", None, "Qty", "Qty"])
        ws.append(["Jan-25", "x", 1, 2])
        ws.append(["Feb-25", "y", 3, 4])

    s = read_workbook_bytes(xlsx_bytes(build), "xlsx").sheets[0]
    assert [c.name for c in s.columns] == ["Month", "Column B", "Qty", "Qty (2)"]


def test_header_falls_back_to_row_1_when_nothing_looks_like_one():
    def build(wb):
        ws = wb.active
        for i in range(4):
            ws.append([i, i * 2, i * 3])

    s = read_workbook_bytes(xlsx_bytes(build), "xlsx").sheets[0]
    assert s.header_row == 1 and s.data_rows == 3


def test_csv_with_semicolons_and_bom():
    content = "﻿Month;Activity;Qty\nJan-25;Diesel;1.200\nFeb-25;Diesel;1.300\n".encode("utf-8")
    s = read_workbook_bytes(content, "csv", "fuel.csv").sheets[0]
    assert [c.name for c in s.columns] == ["Month", "Activity", "Qty"]
    assert s.name == "fuel" and s.row_refs[0] == "fuel!A2"


def test_unknown_extension_is_refused():
    with pytest.raises(ValueError):
        read_workbook_bytes(b"", "pdf")


@pytest.mark.parametrize(
    "header, expected",
    [
        ("Jan", {"month": 1, "year": None, "date": None}),
        ("Jan-24", {"month": 1, "year": 2024, "date": "2024-01-31"}),
        ("January 2025", {"month": 1, "year": 2025, "date": "2025-01-31"}),
        ("Feb'25", {"month": 2, "year": 2025, "date": "2025-02-28"}),
        ("2024-02", {"month": 2, "year": 2024, "date": "2024-02-29"}),
        ("03/2025", {"month": 3, "year": 2025, "date": "2025-03-31"}),
        ("2025-04-01", {"month": 4, "year": 2025, "date": "2025-04-30"}),
        ("Quantity", None),
        ("Mayo", None),
        ("13/2025", None),
    ],
)
def test_month_of_header(header, expected):
    assert month_of_header(header) == expected


# ---------------------------------------------------------------------------
# Review findings (2026-10-14)
# ---------------------------------------------------------------------------


def test_a_sheet_formatted_out_to_the_last_column_reads_quickly():
    import time as _time

    def build(wb):
        ws = wb.active
        ws.append(["Month", "Fuel", "Qty"])
        ws.append(["2025-01", "Diesel", 10])
        ws.append(["2025-02", "Diesel", 12])
        ws["XFD20000"].number_format = "0.00"  # formatted, no value

    content = xlsx_bytes(build)
    t0 = _time.monotonic()
    sheet = read_workbook_bytes(content, "xlsx").sheets[0]
    assert _time.monotonic() - t0 < 10
    assert [c.name for c in sheet.columns] == ["Month", "Fuel", "Qty"] and sheet.data_rows == 2


def test_an_activity_starting_with_total_is_kept_and_a_subtotal_under_a_merged_site_is_dropped():
    def build(wb):
        ws = wb.active
        ws.append(["Site", "Activity", "Qty"])
        ws.append(["Total Energies", "Diesel", 500])
        ws.append(["Plant A", "Diesel", 10])
        ws.append([None, "Petrol", 20])
        ws.append([None, "Subtotal", 30])
        ws.merge_cells("A3:A5")
        ws.append(["Grand total", None, 530])
        ws.append(["Total:", None, 530])

    sheet = read_workbook_bytes(xlsx_bytes(build), "xlsx").sheets[0]
    assert [r["row"] for r in sheet.rows] == [2, 3, 4]
    assert [(d.row, d.reason) for d in sheet.dropped_rows] == [(5, "total"), (6, "total"), (7, "total")]


def test_a_merged_title_above_the_header_is_not_part_of_the_column_names():
    def build(wb):
        ws = wb.active
        ws.append(["Consumption report 2025", None, None])
        ws.merge_cells("A1:C1")
        ws.append(["Month", "Fuel", "Qty"])
        ws.append(["2025-01", "Diesel", 10])
        ws.append(["2025-02", "Diesel", 12])

    sheet = read_workbook_bytes(xlsx_bytes(build), "xlsx").sheets[0]
    assert (sheet.header_row, sheet.header_depth) == (2, 1)
    assert [c.name for c in sheet.columns] == ["Month", "Fuel", "Qty"]


def test_year_headings_are_a_header():
    def build(wb):
        ws = wb.active
        ws.append(["Fuel use by year"])
        ws.append(["Supplier", 2023, 2024, 2025])
        ws.append(["Acme", 100, 110, 120])
        ws.append(["Beta", 200, 210, 220])

    sheet = read_workbook_bytes(xlsx_bytes(build), "xlsx").sheets[0]
    assert sheet.header_row == 2
    assert [c.name for c in sheet.columns] == ["Supplier", "2023", "2024", "2025"] and sheet.data_rows == 2


def test_a_damaged_xlsx_is_a_plain_value_error():
    with pytest.raises(ValueError, match="could not be read"):
        read_workbook_bytes(b"PK\x03\x04garbage", "xlsx")


def test_day_first_date_headers_make_a_months_as_columns_csv():
    header = "Activity," + ",".join(f"01/{m:02d}/2025" for m in range(1, 7))
    content = (header + "\nDiesel,1,2,3,4,5,6\nPetrol,1,2,3,4,5,6\n").encode()
    sheet = read_workbook_bytes(content, "csv", "pivot.csv").sheets[0]
    assert sheet.layout == "months_as_columns"
    assert sheet.columns[1].month == {"month": 1, "year": 2025, "date": "2025-01-31"}
    assert month_of_header("15/03/2025")["month"] == 3


def test_chart_sheets_are_listed_and_keep_their_index():
    def build(wb):
        from openpyxl.chart import BarChart

        a = wb.active
        a.title = "Data1"
        a.append(["Fuel", "Qty"])
        a.append(["Diesel", 1])
        a.append(["Petrol", 2])
        cs = wb.create_chartsheet("Chart1")
        cs.add_chart(BarChart())
        b = wb.create_sheet("Data2")
        b.append(["Fuel", "Qty"])
        b.append(["Diesel", 1])
        b.append(["Petrol", 2])

    sheets = read_workbook_bytes(xlsx_bytes(build), "xlsx").sheets
    assert [(s.name, s.index, s.skipped_reason) for s in sheets] == [("Data1", 0, None), ("Chart1", 1, "chart sheet"), ("Data2", 2, None)]
