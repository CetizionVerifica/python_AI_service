"""KPI-12 · the stored document a suggestion is read from.

Documents live in ``uploaded_documents`` (the same upload storage the bulk
import uses: Cloudinary, or a local temp path on a dev box), by id. A row has
no company of its own, so it belongs to the company of the user who uploaded
it, read from the shared ESG-lite tables the way ``app/core/auth.py`` reads a
caller's sites: the uploader's own site and assigned sites (``"user".site_id``,
``user_sites``) and those sites' ``company_id``. A Superadmin's upload, or one
from before ``uploaded_by`` existed, has no company and is never readable here.

``locate`` answers 404 for a missing document, one of another company, one
uploaded by someone else when ESG-lite names the user, and one whose file was
already consumed: the same answer, so an id never tells another company
anything. ``read_pages`` turns the file into page texts (PDF pages, an image,
spreadsheet sheets, a CSV) for the prompt; nothing here is logged but ids and
counts.
"""
from __future__ import annotations

import csv
import io
import logging
import os
import re
from dataclasses import dataclass
from typing import Optional

from fastapi import HTTPException
from psycopg2.extras import RealDictCursor

from app.core import database
from app.core.config import settings

logger = logging.getLogger(__name__)

PDF_EXTS = {"pdf"}
IMAGE_EXTS = {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"}
SHEET_EXTS = {"xlsx", "xlsm"}
TEXT_EXTS = {"csv", "txt"}
SUPPORTED_EXTS = PDF_EXTS | IMAGE_EXTS | SHEET_EXTS | TEXT_EXTS

PAGE_MAX_CHARS = 12_000
TOTAL_MAX_CHARS = 60_000
SHEET_MAX_ROWS = 400
CONSUMED_STATUSES = {"imported", "deleted"}

NOT_FOUND = HTTPException(status_code=404, detail="Document not found")


@dataclass(frozen=True)
class Page:
    number: int  # 1-based: the page, the sheet, or 1 for a CSV / image
    text: str
    label: str = ""  # "Sheet 'Headcount'" for a spreadsheet; empty otherwise


@dataclass(frozen=True)
class StoredDocument:
    id: int
    name: str
    location: str
    ext: str


def ext_of(name: Optional[str], location: Optional[str] = None) -> str:
    for candidate in (name, location):
        if candidate and "." in candidate:
            ext = candidate.rsplit(".", 1)[-1].lower().split("?")[0]
            if ext:
                return ext
    return ""


def uploader_company_ids(user_id: int) -> set[int]:
    """Companies of the uploader's own sites; empty for a Superadmin or an unknown user."""
    conn = database.get_connection()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute('SELECT role, site_id FROM "user" WHERE user_id = %s', (user_id,))
            user = cur.fetchone()
            if not user or user["role"] == "Superadmin":
                return set()
            cur.execute(
                """
                SELECT DISTINCT s.company_id FROM site s
                WHERE s.company_id IS NOT NULL
                  AND (s.site_id = %s OR s.site_id IN (SELECT site_id FROM user_sites WHERE user_id = %s))
                """,
                (user["site_id"], user_id),
            )
            return {int(r["company_id"]) for r in cur.fetchall()}
    finally:
        database.release_connection(conn)


def locate(document_id: int, company_id: int, user_id: Optional[int] = None) -> StoredDocument:
    """The document, or 404 unless its uploader belongs to ``company_id`` (and is ``user_id`` when given)."""
    doc = database.get_uploaded_document_by_id(document_id)
    if not doc:
        raise NOT_FOUND
    uploader = doc.get("uploaded_by")
    if uploader is None:
        raise NOT_FOUND
    if user_id is not None and int(uploader) != int(user_id):
        raise NOT_FOUND
    if company_id not in uploader_company_ids(int(uploader)):
        raise NOT_FOUND
    if (doc.get("status") or "") in CONSUMED_STATUSES:
        raise NOT_FOUND
    location = doc.get("cloudinary_url") or ""
    name = doc.get("document_name") or f"document-{document_id}"
    return StoredDocument(id=int(doc["id"]), name=name, location=location, ext=ext_of(name, location))


def fetch_bytes(doc: StoredDocument) -> bytes:
    """The stored file: Cloudinary (http[s]) or a local path (dev, legacy rows)."""
    if doc.location.startswith(("http://", "https://")):
        import requests

        try:
            r = requests.get(doc.location, timeout=60)
            r.raise_for_status()
        except Exception as e:
            logger.warning(f"kpi document {doc.id}: storage read failed ({type(e).__name__})")
            raise HTTPException(status_code=502, detail="Could not read the document from storage")
        return r.content
    if doc.location and os.path.isfile(doc.location):
        with open(doc.location, "rb") as f:
            return f.read()
    raise HTTPException(status_code=410, detail="The stored file is no longer available. Please upload it again.")


def _clip(text: str, limit: int = PAGE_MAX_CHARS) -> str:
    text = re.sub(r" {2,}", " ", text or "")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text if len(text) <= limit else text[:limit]


def _pdf_pages(content: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    texts = []
    for page in reader.pages[: settings.MAX_PAGES]:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")
    if sum(len(t.strip()) for t in texts) >= settings.MIN_PDF_TEXT_CHARS:
        return [Page(i + 1, _clip(t)) for i, t in enumerate(texts)]
    return _ocr_pdf_pages(content)


def _ocr_pdf_pages(content: bytes) -> list[Page]:
    """A scanned PDF: render and OCR each page (Tesseract), as the bills pipeline does."""
    import pypdfium2 as pdfium
    import pytesseract

    pdf = pdfium.PdfDocument(content)
    pages = []
    try:
        for i in range(min(settings.MAX_PAGES, len(pdf))):
            image = pdf[i].render(scale=settings.PDF_RENDER_SCALE).to_pil().convert("L")
            pages.append(Page(i + 1, _clip(pytesseract.image_to_string(image))))
    finally:
        pdf.close()
    return pages


def _image_pages(content: bytes) -> list[Page]:
    import pytesseract
    from PIL import Image

    image = Image.open(io.BytesIO(content)).convert("L")
    return [Page(1, _clip(pytesseract.image_to_string(image)))]


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet_pages(content: bytes) -> list[Page]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    pages = []
    try:
        for i, ws in enumerate(wb.worksheets[: settings.MAX_PAGES]):
            lines = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                if any(cells):
                    lines.append("\t".join(cells).rstrip())
                if len(lines) >= SHEET_MAX_ROWS:
                    break
            pages.append(Page(i + 1, _clip("\n".join(lines)), label=f"Sheet '{ws.title}'"))
    finally:
        wb.close()
    return pages


def _text_pages(content: bytes, ext: str) -> list[Page]:
    text = content.decode("utf-8-sig", errors="replace")
    if ext == "csv":
        rows = list(csv.reader(io.StringIO(text)))[:SHEET_MAX_ROWS]
        text = "\n".join("\t".join(c.strip() for c in r) for r in rows if any(c.strip() for c in r))
    return [Page(1, _clip(text))]


def read_pages(doc: StoredDocument, content: bytes) -> list[Page]:
    """Page texts, capped at MAX_PAGES pages and TOTAL_MAX_CHARS characters in all."""
    if doc.ext not in SUPPORTED_EXTS:
        raise HTTPException(status_code=422, detail="This file type can't be read for KPI suggestions")
    try:
        if doc.ext in PDF_EXTS:
            pages = _pdf_pages(content)
        elif doc.ext in IMAGE_EXTS:
            pages = _image_pages(content)
        elif doc.ext in SHEET_EXTS:
            pages = _sheet_pages(content)
        else:
            pages = _text_pages(content, doc.ext)
    except HTTPException:
        raise
    except Exception as e:
        logger.warning(f"kpi document {doc.id}: could not read ({type(e).__name__})")
        raise HTTPException(status_code=422, detail="The document could not be read")
    out, used = [], 0
    for p in pages:
        if used >= TOTAL_MAX_CHARS:
            break
        text = p.text[: TOTAL_MAX_CHARS - used]
        used += len(text)
        out.append(Page(p.number, text, p.label))
    return out
