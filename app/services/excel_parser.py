import io
import logging
import os
import re
from pathlib import Path
from typing import Optional

import openpyxl
import pandas as pd
import requests

from app.services.llm import _call_openrouter, _parse_json_object
from app.schemas.emission_factor import (
    SpreadsheetSchema,
    EmissionFactorRecord,
    ParseExcelResponse,
    DbCategory,
    CategorySuggestion,
    ColumnHeader,
)
from app.core.database import (
    get_uploaded_document_by_id,
    update_uploaded_document,
    list_documents_for_cleanup,
)
from app.core import cloudinary_service

logger = logging.getLogger(__name__)


# ===========================================================================
# PART 1: Emission Factor Excel Parser (LLM-assisted schema detection)
# ===========================================================================

# ---------------------------------------------------------------------------
# Unit normalisation — strip qualifiers like "of material", "of waste" etc.
# "ton of material" → "ton",  "per passenger.km" → "passenger.km"
# ---------------------------------------------------------------------------

_UNIT_STRIP_PATTERNS = [
    # "X of <something>" — keep only X
    re.compile(r"^(.+?)\s+of\s+\w.*$", re.IGNORECASE),
    # leading "per " — drop it
    re.compile(r"^per\s+(.+)$", re.IGNORECASE),
]


_UNIT_ALIASES: dict[str, str] = {
    "tonnes": "tonne",
    "tons": "ton",
    "litres": "litre",
    "liters": "litre",
    "liter": "litre",
    "gallons": "gallon",
    "kilo litre": "kl",
    "kilolitre": "kl",
    "kiloliter": "kl",
    "cubic metre": "cubic meter",
    "m3": "cubic meter",
    "m³": "cubic meter",
    "kilogram": "kg",
    "kilograms": "kg",
    "kgs": "kg",
    "gram": "g",
    "grams": "g",
    "pound": "lb",
    "pounds": "lb",
    "lbs": "lb",
    "meter": "m",
    "meters": "m",
    "metre": "m",
    "metres": "m",
    "kilometer": "km",
    "kilometers": "km",
    "kilometre": "km",
    "kilometres": "km",
    "miles": "mile",
    "mi": "mile",
}


def _normalise_unit(raw: str) -> str:
    """
    Clean a denominator unit extracted from an emission factor sheet.

    Examples:
        "ton of material"   → "ton"
        "Tonnes"            → "tonne"
        "per passenger.km"  → "passenger.km"
        "KWH"              → "kwh"
        "Litre"            → "litre"
    """
    text = raw.strip()
    # Strip parentheses (UOM cells often have wrapping parens like "(kg CO2 e/ km)")
    text = text.replace("(", "").replace(")", "").strip()
    for pat in _UNIT_STRIP_PATTERNS:
        m = pat.match(text)
        if m:
            text = m.group(1).strip()
    lower = text.lower()
    return _UNIT_ALIASES.get(lower, lower)


# Physical unit conversions, kept for the PCF unit reader (``pcf_units``), which
# needs to know how a sheet's "lb", "MWh" or "kg.km" scales to a PCF unit. The
# bulk-upload path no longer converts anything here: ESG-lite's engine does
# that when it plans the rows (IM11).
UNIT_CONVERSIONS = {
    "litre": {"gallon": 0.264172, "ml": 1000, "cubic meter": 0.001, "kilo litre": 0.001, "kl": 0.001},
    "kl": {"litre": 1000, "gallon": 264.172, "ml": 1000000, "cubic meter": 1},
    "kilo litre": {"litre": 1000, "gallon": 264.172, "ml": 1000000, "cubic meter": 1},
    "gallon": {"litre": 3.78541, "ml": 3785.41, "cubic meter": 0.00378541, "kl": 0.00378541},
    "ml": {"litre": 0.001, "gallon": 0.000264172, "kl": 0.000001},
    "cubic meter": {"litre": 1000, "gallon": 264.172, "kl": 1},
    # Weight
    "kg": {"lb": 2.20462, "tonne": 0.001, "g": 1000, "ton": 0.00110231},
    "lb": {"kg": 0.453592, "tonne": 0.000453592, "g": 453.592},
    "ton": {"kg": 907.185, "lb": 2000, "g": 907185},
    "tonne": {"kg": 1000, "lb": 2204.62, "g": 1000000},
    "g": {"kg": 0.001, "lb": 0.00220462},
    # Freight (weight-distance), the Node backend's table
    "tonne.km": {"kg.km": 1000, "g.km": 1000000},
    "kg.km": {"tonne.km": 0.001, "g.km": 1000},
    "g.km": {"tonne.km": 0.000001, "kg.km": 0.001},
    # Energy
    "kwh": {"mwh": 0.001, "gj": 0.0036, "mj": 3.6},
    "mwh": {"kwh": 1000, "gj": 3.6, "mj": 3600},
    "gj": {"kwh": 277.778, "mwh": 0.277778, "mj": 1000},
    "mj": {"kwh": 0.277778, "gj": 0.001},
}


# ---------------------------------------------------------------------------
# System prompt for Pass 1: LLM schema detection
# ---------------------------------------------------------------------------

SCHEMA_DETECTION_PROMPT = """\
You are an expert at analyzing Excel spreadsheet structures for emission factor data.

You will receive the first ~20 rows of an Excel spreadsheet, represented as a list of
rows where each row is a list of (cell_value, column_index) tuples for non-empty cells.
Column indices are 1-based (A=1, B=2, ...).

Your task is to identify the STRUCTURE of this spreadsheet so that a program can
deterministically extract all data rows. Analyze the header rows to determine:

1. LAYOUT TYPE — one of three patterns:
   - "simple": One factor value per year per row.
     Each year has exactly 1 value column.
   - "sub_columns": Multiple sub-columns per year (e.g. Direct, WTT, Total).
     Each year repeats the same set of sub-columns. The program will pick one
     sub-column (typically "Total") as the factor value.
   - "disposal_pivot": Multiple disposal/treatment method columns per year
     (e.g. Re-use, Open loop, Closed loop, Combustion, Composting, Landfilled,
     Anaerobically digested). Each disposal column × each data row = a separate
     emission factor record.

2. DESCRIPTOR COLUMNS — which columns contain text descriptors that together form
   the emission_category_name. These are text columns to the LEFT of the numeric
   factor data. Report them in left-to-right order. Exclude:
   - Broad "Category" or "Method of calculation" columns that are the same for all rows
   - Source / UOM columns
   Only include columns whose values differentiate emission factor records.

3. PARENT CATEGORY COLUMN — an optional grouping column (e.g. "Stationary Combustion",
   "Mobile Combustion", "Road", "Air"). Rows with a value ONLY in this column and no
   numeric data are group headers. Set `include_parent_in_name` to true ONLY if the
   same descriptor values appear under different parents and need disambiguation
   (e.g. "Diesel" under both "Stationary Combustion" and "Mobile Combustion").

4. UNIT COLUMN — which column (if any) contains the unit of measurement (UOM).

5. SOURCE COLUMN — which column (if any) contains the data source (e.g. "DEFRA").

6. DATA START ROW — the first row number containing actual data values (not headers).

7. YEAR COLUMNS — for each year detected:
   - "simple": provide `value_column` (the column index of the FACTOR VALUE, NOT the year label).
   - "sub_columns": provide `sub_columns` (list of {name, column_index}) and
     `primary_sub_column` (name of the sub-column to use, usually "Total").
   - "disposal_pivot": provide `disposal_columns` (list of {name, column_index}).

CRITICAL — FACTOR VALUE vs YEAR LABEL DISAMBIGUATION:
   Sometimes a sheet has a single numeric factor value column (e.g. header "Emission factors",
   "Factor", "Value") AND separate year columns that just list which years the factor applies to.
   For example:
     | Ef Category | Unit | Emission factors | year |      |      |      |
     | Diesel      | USD  | 0.134            | 2025 | 2024 | 2023 | 2022 |
   Here column 3 ("Emission factors") holds the actual factor value (0.134).
   Columns 4-7 hold year LABELS (2025, 2024, 2023, 2022) — they indicate this
   factor applies to all those years, but the VALUE is the SAME for each year (0.134).
   In this case you MUST set `value_column` to the factor value column (3), NOT
   to the year label columns (4-7). Each year entry should share the same value_column.
   Key signal: if the data cells in columns 4+ are year-like integers (2015-2035)
   and a separate column has small decimal values, those integers are year labels,
   not factor values.

Return ONLY a JSON object with this exact structure (no extra text):
{
  "layout_type": "simple" | "sub_columns" | "disposal_pivot",
  "descriptor_columns": [
    {"column_index": <int>, "header_name": "<string>"}
  ],
  "parent_category_column": {"column_index": <int>, "header_name": "<string>"} | null,
  "include_parent_in_name": true | false,
  "source_column": {"column_index": <int>, "header_name": "<string>"} | null,
  "unit_column": {"column_index": <int>, "header_name": "<string>"} | null,
  "data_start_row": <int>,
  "years": [
    {
      "year": <int>,
      "value_column": <int or null>,
      "sub_columns": [{"name": "<string>", "column_index": <int>}] | null,
      "primary_sub_column": "<string>" | null,
      "disposal_columns": [{"name": "<string>", "column_index": <int>}] | null
    }
  ],
  "descriptor_join_separator": " - ",
  "notes": "<any observations about the data>"
}
"""


# ---------------------------------------------------------------------------
# Pass 1 helpers
# ---------------------------------------------------------------------------

def _read_header_rows(ws, max_rows: int = 20) -> str:
    """Read first N rows of worksheet, format as (value, column_index) tuples."""
    lines: list[str] = []
    for row in ws.iter_rows(
        min_row=1, max_row=min(max_rows, ws.max_row), values_only=False
    ):
        row_num = row[0].row
        cells = [
            (cell.value, cell.column)
            for cell in row
            if cell.value is not None
        ]
        if cells:
            lines.append(f"Row {row_num}: {cells}")
    return "\n".join(lines)


def _detect_schema(header_text: str) -> SpreadsheetSchema:
    """Pass 1: Send header rows to LLM for structure detection."""
    messages = [
        {"role": "system", "content": SCHEMA_DETECTION_PROMPT},
        {"role": "user", "content": f"SPREADSHEET ROWS:\n{header_text}"},
    ]
    logger.info("Pass 1: Calling LLM for schema detection")
    raw_response = _call_openrouter(messages)
    schema_dict = _parse_json_object(raw_response)
    logger.info(f"Pass 1: LLM returned schema: layout_type={schema_dict.get('layout_type')}")
    return SpreadsheetSchema(**schema_dict)


# ---------------------------------------------------------------------------
# Pass 2: Deterministic extraction
# ---------------------------------------------------------------------------

def _parse_numeric(value, precision: int = 6) -> Optional[float]:
    """Safely convert a cell value to float, returning None for non-numeric.

    Rounds to *precision* decimal places to strip floating-point noise
    that arises from Excel formula caching (e.g. 0.74994 stored as
    0.7499399999999999).
    """
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, bool):
            return None
        return round(float(value), precision)
    s = str(value).strip()
    if s in ("", "-", "N/A", "n/a", "NA", "na", "#N/A", "#REF!", "#VALUE!"):
        return None
    try:
        return round(float(s), precision)
    except ValueError:
        return None


# emission_factors.factor_value is numeric(10,4): four decimal places are all
# the database keeps. Factors are rounded to that here, so what the review
# screen shows is exactly what gets stored, and a factor that loses digits is
# flagged instead of silently truncated.
FACTOR_DECIMALS = 4


def _parse_factor(value, row_num: int, name: str, warnings: list[str]) -> Optional[float]:
    parsed = _parse_numeric(value)
    if parsed is None:
        return None
    stored = round(parsed, FACTOR_DECIMALS)
    if abs(stored - parsed) > 1e-9:
        warnings.append(
            f"Row {row_num}: factor {parsed:g} for '{name}' is stored as {stored:.{FACTOR_DECIMALS}f} "
            f"(the database keeps {FACTOR_DECIMALS} decimal places)"
        )
    return stored


def _extract_factors(
    ws, schema: SpreadsheetSchema
) -> tuple[list[EmissionFactorRecord], list[str]]:
    """
    Pass 2: Use detected schema to programmatically read all data rows.
    No LLM involved — purely deterministic openpyxl reads.
    """
    factors: list[EmissionFactorRecord] = []
    warnings: list[str] = []

    current_parent: Optional[str] = None
    last_unit: Optional[str] = None
    last_source: Optional[str] = None
    # Forward-fill for merged descriptor cells: column_index -> last non-empty value
    last_descriptors: dict[int, str] = {}

    for row in ws.iter_rows(
        min_row=schema.data_start_row, max_row=ws.max_row, values_only=False
    ):
        row_num = row[0].row
        # Build a column_index -> cell_value lookup
        cells: dict[int, any] = {cell.column: cell.value for cell in row}

        # --- Track parent category (for group-header rows) ---
        if schema.parent_category_column:
            parent_val = cells.get(schema.parent_category_column.column_index)
            if parent_val is not None and str(parent_val).strip():
                current_parent = str(parent_val).strip()

        # --- Build emission_category_name from descriptors ---
        # Forward-fill handles merged cells: when a cell is empty, use the
        # last value for that column.  When a cell has a new value, reset
        # all columns to its right so stale child-level values don't leak.
        descriptor_parts: list[str] = []
        if schema.include_parent_in_name and current_parent:
            descriptor_parts.append(current_parent)

        has_any_explicit = False  # track if row has at least one real cell
        for i, desc_col in enumerate(schema.descriptor_columns):
            val = cells.get(desc_col.column_index)
            val_str = str(val).strip() if val is not None else ""

            if val_str and val_str != "-":
                # Explicit value — update forward-fill and reset downstream
                has_any_explicit = True
                last_descriptors[desc_col.column_index] = val_str
                for j in range(i + 1, len(schema.descriptor_columns)):
                    last_descriptors.pop(
                        schema.descriptor_columns[j].column_index, None
                    )
                descriptor_parts.append(val_str)
            elif desc_col.column_index in last_descriptors:
                # Empty / merged cell — use forward-filled value
                descriptor_parts.append(last_descriptors[desc_col.column_index])
            # else: truly empty and no previous value, skip

        if not descriptor_parts or (
            not has_any_explicit
            and not schema.include_parent_in_name
        ):
            continue  # Skip empty / group-header rows with no descriptors

        emission_category_name = schema.descriptor_join_separator.join(
            descriptor_parts
        )

        # --- Source and unit ---
        source: Optional[str] = None
        if schema.source_column:
            src_val = cells.get(schema.source_column.column_index)
            if src_val is not None and str(src_val).strip():
                source = str(src_val).strip()
                last_source = source
            elif last_source:
                source = last_source

        denominator_unit: Optional[str] = None
        if schema.unit_column:
            unit_val = cells.get(schema.unit_column.column_index)
            if unit_val is not None and str(unit_val).strip():
                raw_unit = str(unit_val).strip()
                # Extract just the denominator (e.g. "KgCO2e/USD" → "USD")
                denominator_unit = raw_unit.split("/", 1)[1].strip() if "/" in raw_unit else raw_unit
                # Normalise: "ton of material" → "ton", "per passenger.km" → "passenger.km"
                denominator_unit = _normalise_unit(denominator_unit)
                last_unit = denominator_unit
            elif last_unit:
                denominator_unit = last_unit
                warnings.append(
                    f"Row {row_num}: Empty unit cell, inherited '{last_unit}' from previous row"
                )

        # --- Extract factor values based on layout type ---
        row_has_data = False

        for year_map in schema.years:
            if schema.layout_type == "simple":
                if year_map.value_column is None:
                    continue
                value = cells.get(year_map.value_column)
                factor = _parse_factor(value, row_num, emission_category_name, warnings)
                if factor is not None:
                    row_has_data = True
                    factors.append(
                        EmissionFactorRecord(
                            year=year_map.year,
                            factor_value=factor,
                            denominator_unit=denominator_unit,
                            source=source,
                            emission_category_name=emission_category_name,
                            parent_category=current_parent,
                        )
                    )

            elif schema.layout_type == "sub_columns":
                # Find the primary sub-column (e.g., "Total")
                total_col: Optional[int] = None
                for sc in year_map.sub_columns or []:
                    if sc.name == year_map.primary_sub_column:
                        total_col = sc.column_index
                        break
                if total_col is None:
                    continue
                value = cells.get(total_col)
                factor = _parse_factor(value, row_num, emission_category_name, warnings)
                if factor is not None:
                    row_has_data = True
                    factors.append(
                        EmissionFactorRecord(
                            year=year_map.year,
                            factor_value=factor,
                            denominator_unit=denominator_unit,
                            source=source,
                            emission_category_name=emission_category_name,
                            parent_category=current_parent,
                        )
                    )

            elif schema.layout_type == "disposal_pivot":
                # Each disposal column produces a separate record
                for disp_col in year_map.disposal_columns or []:
                    value = cells.get(disp_col.column_index)
                    pivot_name = (
                        f"{emission_category_name}"
                        f"{schema.descriptor_join_separator}"
                        f"{disp_col.name}"
                    )
                    factor = _parse_factor(value, row_num, pivot_name, warnings)
                    if factor is not None:
                        row_has_data = True
                        factors.append(
                            EmissionFactorRecord(
                                year=year_map.year,
                                factor_value=factor,
                                denominator_unit=denominator_unit,
                                source=source,
                                emission_category_name=pivot_name,
                                parent_category=current_parent,
                            )
                        )

        if not row_has_data and descriptor_parts:
            # This might be a group-header row — not a warning
            pass

    return factors, warnings


# ---------------------------------------------------------------------------
# Post-processing: disambiguate duplicate names with different units
# ---------------------------------------------------------------------------

def _disambiguate_by_unit(
    factors: list[EmissionFactorRecord],
) -> list[EmissionFactorRecord]:
    """
    When the same emission_category_name appears with different denominator_units
    (e.g. "Road - Van - Diesel" at both tonne.km and km), append a unit-based
    suffix to make them unique.
    """
    from collections import defaultdict

    name_units: dict[str, set[str]] = defaultdict(set)
    for f in factors:
        if f.emission_category_name and f.denominator_unit:
            name_units[f.emission_category_name].add(f.denominator_unit)

    # Only names that appear with 2+ different units need disambiguation
    ambiguous = {name for name, units in name_units.items() if len(units) > 1}
    if not ambiguous:
        return factors

    def _unit_suffix(unit: str) -> str:
        """Extract the meaningful denominator part: 'Kg CO2e/tonne.km' → 'tonne.km'"""
        if "/" in unit:
            return unit.split("/", 1)[1].strip()
        return unit

    result: list[EmissionFactorRecord] = []
    for f in factors:
        if f.emission_category_name in ambiguous and f.denominator_unit:
            suffix = _unit_suffix(f.denominator_unit)
            result.append(
                f.model_copy(
                    update={
                        "emission_category_name": f"{f.emission_category_name} [{suffix}]"
                    }
                )
            )
        else:
            result.append(f)
    return result


# ---------------------------------------------------------------------------
# Column header extraction — for user-facing column mapping UI
# ---------------------------------------------------------------------------

def _read_available_columns(ws, max_sample_rows: int = 5) -> list[ColumnHeader]:
    """Read all column headers from the worksheet with a few sample values."""
    columns: list[ColumnHeader] = []

    # Read header row (row 1)
    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=False))[0]

    for cell in header_row:
        header_name = str(cell.value).strip() if cell.value is not None else f"Column {cell.column}"
        # Gather sample values from next few rows
        samples: list[str] = []
        for data_row in ws.iter_rows(
            min_row=2,
            max_row=min(1 + max_sample_rows, ws.max_row),
            min_col=cell.column,
            max_col=cell.column,
            values_only=True,
        ):
            val = data_row[0]
            if val is not None:
                samples.append(str(val)[:80])  # truncate long values
        columns.append(ColumnHeader(
            column_index=cell.column,
            header_name=header_name,
            sample_values=samples,
        ))

    return columns


# ---------------------------------------------------------------------------
# Validation — catch factor values that look like year numbers
# ---------------------------------------------------------------------------

def _validate_factors(
    factors: list[EmissionFactorRecord],
) -> list[str]:
    """
    Post-extraction sanity checks. Returns warnings for suspicious data.
    Catches the bug where year numbers (2020-2030) are read as factor values.
    """
    warnings: list[str] = []
    current_year = 2026
    year_range = set(range(2015, current_year + 10))

    year_value_count = sum(
        1 for f in factors if int(f.factor_value) == f.factor_value and int(f.factor_value) in year_range
    )
    total = len(factors)

    if total > 0 and year_value_count / total > 0.5:
        warnings.append(
            f"DATA QUALITY WARNING: {year_value_count}/{total} factor values "
            f"look like year numbers (e.g. 2022, 2023). The column mapping may "
            f"be incorrect — the system might be reading year labels as factor "
            f"values. Please verify the 'Factor Value' column mapping."
        )

    return warnings


# ---------------------------------------------------------------------------
# Public entry point — Emission Factor Excel Parser
# ---------------------------------------------------------------------------

def parse_emission_factor_excel(
    file_path: Path, filename: str, sheet_name: Optional[str] = None,
    schema_override: Optional[dict] = None,
) -> ParseExcelResponse:
    """Main entry: two-pass Excel parsing.

    Args:
        file_path: Path to the Excel file.
        filename: Original filename (for logging / response).
        sheet_name: Specific sheet to parse. If None, uses the first sheet.
        schema_override: If provided, skip LLM detection and use this schema
                         directly for extraction. Allows user column mapping overrides.
    """
    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    sheet_names = wb.sheetnames

    # Select worksheet
    if sheet_name and sheet_name in sheet_names:
        ws = wb[sheet_name]
        selected_sheet = sheet_name
    else:
        ws = wb[sheet_names[0]]
        selected_sheet = sheet_names[0]

    # Read available columns for the mapping UI
    available_columns = _read_available_columns(ws)

    if schema_override:
        # User provided column mapping — skip LLM, extract directly
        logger.info(f"Using user-provided schema override for {filename}")
        schema = SpreadsheetSchema(**schema_override)
    else:
        # Pass 1: Schema detection via LLM
        header_text = _read_header_rows(ws, max_rows=20)
        logger.info(f"Pass 1: Detecting schema for {filename}")
        schema = _detect_schema(header_text)
        logger.info(
            f"Detected layout: {schema.layout_type}, "
            f"years: {[y.year for y in schema.years]}, "
            f"descriptors: {[d.header_name for d in schema.descriptor_columns]}"
        )

    # Pass 2: Deterministic extraction
    logger.info(f"Pass 2: Extracting factors from {filename} (sheet: {selected_sheet})")
    factors, warnings = _extract_factors(ws, schema)

    wb.close()

    # Pass 3: Disambiguate duplicate names that differ only by unit
    factors = _disambiguate_by_unit(factors)

    # Pass 4: Validate extracted data for suspicious patterns
    validation_warnings = _validate_factors(factors)
    warnings.extend(validation_warnings)

    available_years = sorted(set(f.year for f in factors))
    parent_categories = sorted(
        set(f.parent_category for f in factors if f.parent_category)
    )

    logger.info(
        f"Extracted {len(factors)} emission factor records "
        f"across years {available_years} from {filename} (sheet: {selected_sheet})"
    )

    return ParseExcelResponse(
        filename=filename,
        factors=factors,
        schema_detected=schema,
        warnings=warnings,
        total_records=len(factors),
        available_years=available_years,
        parent_categories=parent_categories,
        sheet_names=sheet_names,
        selected_sheet=selected_sheet,
        available_columns=available_columns,
    )


# ---------------------------------------------------------------------------
# Category inference via LLM
# ---------------------------------------------------------------------------

CATEGORY_INFERENCE_PROMPT = """\
You are an expert at matching emission factor categories.

You will receive:
1. A list of PARENT CATEGORIES extracted from an Excel emission factor spreadsheet.
2. A list of DATABASE CATEGORIES (each with an id and name) from the user's system.

Your task is to match each parent category to the most appropriate database category.
Consider semantic similarity, not just exact string matching. For example:
- "Stationary Combustion" might match "Scope 1 - Stationary Combustion"
- "Business travel" might match "Business Travel" or "Scope 3 - Business Travel"
- "Downstream transportation and distribustion" (note typo) should still match "Downstream Transportation"

For each parent category, assign a confidence level:
- "high": Very clear match (near-exact or obvious semantic equivalence)
- "medium": Reasonable match but some ambiguity
- "low": No good match found

If no database category is a reasonable match, set suggested_category_id to null.

Return ONLY a JSON object (no extra text):
{
  "suggestions": [
    {
      "parent_category": "<exact parent category string>",
      "suggested_category_id": <int or null>,
      "suggested_category_name": "<name of matched DB category or null>",
      "confidence": "high" | "medium" | "low"
    }
  ]
}
"""


def infer_category_mapping(
    parent_categories: list[str],
    db_categories: list[DbCategory],
) -> list[CategorySuggestion]:
    """Use LLM to match Excel parent categories to DB categories."""
    user_content = (
        f"PARENT CATEGORIES FROM EXCEL:\n"
        f"{parent_categories}\n\n"
        f"DATABASE CATEGORIES:\n"
        f"{[{'id': c.id, 'name': c.name} for c in db_categories]}"
    )

    messages = [
        {"role": "system", "content": CATEGORY_INFERENCE_PROMPT},
        {"role": "user", "content": user_content},
    ]

    logger.info(
        f"Category inference: matching {len(parent_categories)} parent categories "
        f"against {len(db_categories)} DB categories"
    )
    raw_response = _call_openrouter(messages)
    result = _parse_json_object(raw_response)

    suggestions = [
        CategorySuggestion(**item) for item in result.get("suggestions", [])
    ]

    logger.info(
        f"Category inference complete: "
        f"{sum(1 for s in suggestions if s.suggested_category_id is not None)} matched"
    )
    return suggestions


# ===========================================================================
# PART 2: Data Import Excel Parser (Pavithra's bulk upload flow)
# ===========================================================================

# ---------------------------------------------------------------------------
# In-memory file cache — avoids re-downloading from Cloudinary on each step
# TTL = 10 minutes, auto-evicts stale entries
# ---------------------------------------------------------------------------
import time as _time

_file_cache: dict[int, tuple[bytes, str, float]] = {}  # doc_id → (bytes, ext, timestamp)
_FILE_CACHE_TTL = 600  # 10 minutes


def _cache_get(document_id: int) -> tuple[bytes, str] | None:
    entry = _file_cache.get(document_id)
    if entry is None:
        return None
    content, ext, ts = entry
    if _time.monotonic() - ts > _FILE_CACHE_TTL:
        _file_cache.pop(document_id, None)
        return None
    return content, ext


def _cache_set(document_id: int, content: bytes, ext: str) -> None:
    # Evict stale entries (keep cache bounded)
    now = _time.monotonic()
    stale = [k for k, (_, _, ts) in _file_cache.items() if now - ts > _FILE_CACHE_TTL]
    for k in stale:
        _file_cache.pop(k, None)
    _file_cache[document_id] = (content, ext, now)

def _cache_delete(document_id: int) -> None:
    _file_cache.pop(document_id, None)


# ---------- Excel reading ----------
# How many leading rows to consider as candidate header rows. Real client
# workbooks often carry a title banner, a blank spacer, and sometimes a row of
# column-letter labels before the actual headers.
_HEADER_SCAN_ROWS = 10

# A row whose non-empty cells are all short uppercase letter codes ("A", "B",
# "AC") is a column-letter legend, not data. Glochem's sheets carry one directly
# under the header row.
_LETTER_LABEL_RE = re.compile(r"^[A-Z]{1,2}$")


def _clean_header(name) -> str:
    """
    Collapse a header cell to a single-line, single-spaced string.

    Excel headers routinely wrap ("Received\\nDate", "UOM\\nINR/USD"). The mapping
    screen and the import must agree on the exact spelling, so both go through here.
    """
    return re.sub(r"\s+", " ", str(name)).strip()


def _named_columns(df: pd.DataFrame) -> list[str]:
    """Columns that came from a real header cell (not pandas' Unnamed: N filler)."""
    return [
        c for c in df.columns
        if not str(c).startswith("Unnamed:") and str(c).strip() and str(c).strip().lower() != "nan"
    ]


def _read_raw(content: bytes, ext: str, header_row: int, nrows: int | None = None) -> pd.DataFrame:
    if ext == "csv":
        return pd.read_csv(io.BytesIO(content), dtype=str, header=header_row, nrows=nrows)
    return pd.read_excel(io.BytesIO(content), dtype=str, header=header_row, nrows=nrows)


def detect_header_row(content: bytes, ext: str) -> int:
    """
    Pick the row that best looks like the header row.

    The previous logic took the *first* row that produced any named column, which
    on a workbook with a merged title banner returns that banner as a lone bogus
    header. Scoring every candidate and keeping the richest one handles banners,
    blank spacers and multi-line titles alike.
    """
    best_row, best_score = 0, -1
    for header_row in range(_HEADER_SCAN_ROWS):
        try:
            df = _read_raw(content, ext, header_row, nrows=5)
        except Exception:
            continue
        named = _named_columns(df)
        # A single named column is almost always a merged title cell.
        if len(named) < 2:
            continue
        if len(named) > best_score:
            best_score, best_row = len(named), header_row
    if best_score < 0:
        raise ValueError("Failed to parse file: no valid headers found.")
    return best_row


def _drop_legend_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Drop column-letter legend rows (e.g. 'A | B | C | D') sitting under the header."""
    def is_legend(row) -> bool:
        vals = [str(v).strip() for v in row.tolist() if str(v).strip()]
        return len(vals) >= 3 and all(_LETTER_LABEL_RE.match(v) for v in vals)

    keep = ~df.apply(is_legend, axis=1)
    dropped = int((~keep).sum())
    if dropped:
        logger.info(f"Dropped {dropped} column-letter legend row(s) below the header.")
    return df[keep]


def read_headers_from_bytes(content: bytes, ext: str) -> list[str]:
    """Header names as the import will see them — same detection as _read_df_from_bytes."""
    header_row = detect_header_row(content, ext)
    df = _read_raw(content, ext, header_row, nrows=5)
    return [_clean_header(c) for c in _named_columns(df)]


def _read_df_from_bytes(content: bytes, ext: str) -> pd.DataFrame:
    header_row = detect_header_row(content, ext)
    df = _read_raw(content, ext, header_row)
    df = df.fillna("")
    df = _drop_legend_rows(df)
    # Match the spelling handed to the mapping screen by read_headers_from_bytes.
    df.columns = [_clean_header(c) for c in df.columns]
    logger.info(f"Parsed sheet using header row {header_row + 1}; {len(df)} data rows.")
    return df

def map_df(df: pd.DataFrame, mappings: dict[str, str]) -> pd.DataFrame:
    mapped = {}
    for required_field, excel_col in mappings.items():
        col = str(excel_col).strip()
        if col in df.columns:
            mapped[required_field] = df[col]
        else:
            logger.warning(f"Column '{col}' not found in file, skipping.")
    if not mapped:
        raise ValueError("No valid column mappings found.")
    return pd.DataFrame(mapped).fillna("")

def _resolve_ext(document_name: str | None, location: str) -> str:
    """
    Work out which spreadsheet reader to use for a stored document.

    The original filename is checked first: a Cloudinary raw URL does not
    reliably carry a usable extension, and silently defaulting to xlsx would
    mis-parse a CSV.
    """
    for candidate in (document_name, location.split("?")[0]):
        if not candidate or "." not in candidate:
            continue
        ext = candidate.rsplit(".", 1)[-1].strip().lower()
        if ext in {"csv", "xls", "xlsx"}:
            return ext
    return "xlsx"


def download_document_bytes(document_id: int) -> tuple[bytes, str]:
    # Check in-memory cache first (avoids re-reading on each step)
    cached = _cache_get(document_id)
    if cached is not None:
        return cached

    doc = get_uploaded_document_by_id(document_id)
    if not doc:
        raise ValueError("Invalid document_id")

    # The file is deleted the moment an import commits, so a request for an
    # already-consumed document is expected rather than exceptional. Check the
    # status first: Cloudinary's CDN can keep serving a deleted asset for a
    # while, and reading it back after import would be silently inconsistent.
    if (doc.get("status") or "") in CONSUMED_STATUSES:
        raise ValueError(
            "This file has already been imported and is no longer stored. "
            "Please re-upload it to import again."
        )

    file_path = doc["cloudinary_url"]
    ext = _resolve_ext(doc.get("document_name"), file_path)

    # Shared storage (Cloudinary). This is the normal path: any instance can
    # fetch it, and it survives restarts. The bytes are cached below so paging
    # through the preview does not re-download the file each time.
    if file_path.startswith(("http://", "https://")):
        try:
            r = requests.get(file_path, timeout=60)
            r.raise_for_status()
        except Exception as e:
            raise ValueError(f"Failed to read file: {e}")

        _cache_set(document_id, r.content, ext)
        return r.content, ext

    # Legacy rows still point at a container-local path.
    if os.path.isfile(file_path):
        with open(file_path, "rb") as f:
            content = f.read()
        _cache_set(document_id, content, ext)
        return content, ext

    raise ValueError(
        f"Temp file no longer exists at '{file_path}'. Please re-upload the file."
    )

# ---------- Per-row reporting dates ----------
# A client export typically holds a full year of transactions with a real date on
# every line. Mapping that column lets one upload land in the right month (and the
# right year-lagged emission factor) per row, instead of stamping the whole file
# with a single import date.

_DATE_FIELD = "date_of_reporting"


def _parse_row_date(value, fallback: str) -> str:
    """
    Normalise a spreadsheet date cell to YYYY-MM-DD.

    Day-first is assumed for ambiguous numeric dates: these are Indian client
    exports where 03.04.2025 means 3 April, not 4 March. Falls back to the
    import-level date when the cell is blank or unparseable.
    """
    raw = str(value or "").strip()
    if not raw:
        return fallback
    try:
        ts = pd.to_datetime(raw, dayfirst=True, errors="coerce")
        if pd.isna(ts):
            return fallback
        return ts.strftime("%Y-%m-%d")
    except Exception:
        return fallback


def _row_reporting_date(activity_data: dict, fallback: str) -> str:
    """Per-row date if a date column was mapped, otherwise the import-level date."""
    if _DATE_FIELD in activity_data:
        return _parse_row_date(activity_data.get(_DATE_FIELD), fallback)
    return fallback


# ---------- Step 2: unique categories ----------
def get_unique_categories(document_id: int, mappings: dict[str, str]) -> tuple[list[str], int]:
    content, ext = download_document_bytes(document_id)
    df = _read_df_from_bytes(content, ext)
    mapped_df = map_df(df, mappings)

    if "emission_category" not in mapped_df.columns:
        raise ValueError("Mapping must include emission_category")

    cats = sorted(
        {str(x).strip() for x in mapped_df["emission_category"].tolist() if str(x).strip()}
    )
    return cats, len(mapped_df)


# Statuses that mean the stored file is gone on purpose. download_document_bytes
# uses these to explain *why* a document can no longer be read.
CONSUMED_STATUSES = {"imported", "deleted"}


def delete_document_file(document_id: int, new_status: str = "deleted") -> None:
    """
    Permanently remove a document's stored file (Cloudinary asset or legacy
    local temp file) and drop it from the in-memory cache.

    Called as soon as an import commits: the rows are in the database, so the
    spreadsheet has served its purpose. A retry means re-uploading the file.
    """
    try:
        doc = get_uploaded_document_by_id(document_id)
        if not doc:
            return

        public_id = doc.get("cloudinary_public_id")
        file_path = doc.get("cloudinary_url") or ""

        if public_id:
            cloudinary_service.delete_file(public_id)
            logger.info(f"Deleted Cloudinary asset for document_id={document_id}: {public_id}")
        elif file_path and os.path.isfile(file_path):
            os.remove(file_path)
            logger.info(f"Deleted temp file for document_id={document_id}: {file_path}")

        update_uploaded_document(document_id, status=new_status)
    except Exception as e:
        logger.warning(f"Failed to delete stored file for document_id={document_id}: {type(e).__name__}")
        # After an import the rows are saved whatever happened to the file, so
        # the document still moves on from "importing" (else a retry is refused
        # as "already being imported" for a batch that landed) and the file is
        # left for cleanup_old_documents. A cleanup that fails keeps its status,
        # so it is retried on the next run.
        if new_status == "imported":
            try:
                update_uploaded_document(document_id, status=new_status)
            except Exception as e2:
                logger.error(f"Could not update status for document_id={document_id}: {type(e2).__name__}")
    finally:
        _cache_delete(document_id)


def cleanup_old_documents(older_than_days: int = 7, limit: int = 500) -> int:
    """
    Reclaim storage for *abandoned* uploads — files whose wizard was started but
    never imported, so nothing ever deleted them. Imported documents clean
    themselves up the moment their rows commit.

    Optional; nothing calls it. Returns the number of documents cleaned up.
    """
    try:
        docs = list_documents_for_cleanup(older_than_days=older_than_days, limit=limit)
    except Exception as e:
        logger.error(f"cleanup_old_documents: failed to list documents: {e}")
        return 0

    for doc in docs:
        delete_document_file(int(doc["id"]))

    logger.info(f"cleanup_old_documents: processed {len(docs)} document(s)")
    return len(docs)


# ---------- The import path (IM11) ----------
# The calculation and the emission rows are ESG-lite's since IM11: /v1/excel/preview
# and /import hand the mapped rows to its plans API (app/services/excel_plan_client.py).
# This module only reads the sheet the way the mapping screen saw it.


def load_mapped_sheet(document_id: int, mappings: dict[str, str]) -> tuple[pd.DataFrame, str | None]:
    """The stored sheet with the mapped columns only, plus the file's name for the plan's source_name."""
    content, ext = download_document_bytes(document_id)
    df = _read_df_from_bytes(content, ext)
    mapped_df = map_df(df, mappings)
    doc = get_uploaded_document_by_id(document_id) or {}
    return mapped_df, (doc.get("document_name") or None)
