"""PCF material-factor sheet reader (E2): any layout → MaterialFactor rows.

Reads the first useful worksheet (or the one asked for), finds the header row,
matches columns to MaterialFactor fields by name, and asks the LLM only for the
required columns names alone can't place. Every row comes back with the
MaterialFactor body ESG-lite's ``POST /pcf/material-factors/import`` takes, a
``confidence`` (0–100) and a ``reason``; anything the reader is unsure of is a
warning and anything ESG-lite would refuse is a problem. Nothing is saved here:
the person checks the rows in C04 and ESG-lite imports them.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import logging
import re
from dataclasses import dataclass, field
from typing import Any, Callable, Optional

import openpyxl
from rapidfuzz import fuzz

from app.services.pcf_units import PcfUnit, co2e_multiplier, to_pcf_unit

logger = logging.getLogger(__name__)

MAX_ROWS = 5000
HEADER_SCAN_ROWS = 15
SAMPLE_ROWS = 5

GROUPS = ("aluminium", "copper", "polymer", "steel", "packaging", "chemical", "energy", "transport", "waste")

# (field, label, required). The order is the matching order: a header is used once.
FIELDS: list[tuple[str, str, bool]] = [
    ("value_kgco2e", "Value (kgCO2e per unit)", True),
    ("name", "Name", True),
    ("material_group", "Group", True),
    ("unit", "Unit", True),
    ("geography", "Geography", False),
    ("gwp_set", "GWP set", False),
    ("source_year", "Source year", False),
    ("source", "Source", False),
    ("dataset_ref", "Dataset ref", False),
    ("licence", "Licence", False),
    ("recycled_variant", "Recycled variant", False),
    ("valid_from", "Valid from", False),
    ("valid_to", "Valid to", False),
]
FIELD_LABEL = {f: label for f, label, _ in FIELDS}
# Group and unit can also come from the name and the value header.
NEEDS_A_COLUMN = ("name", "value_kgco2e")

SYNONYMS: dict[str, list[str]] = {
    "name": ["name", "material", "materialname", "factorname", "description", "item", "product", "activity", "process", "datasetname", "materialdescription"],
    "material_group": ["materialgroup", "group", "category", "materialcategory", "type", "materialtype", "class"],
    "unit": ["unit", "units", "declaredunit", "perunit", "uom", "unitofmeasure", "referenceunit", "functionalunit"],
    "value_kgco2e": ["valuekgco2e", "value", "kgco2e", "factor", "factorvalue", "kgco2eperunit", "gwp", "gwp100", "emissionfactor", "co2e", "carbonfootprint", "pcf", "gwptotal", "climatechange", "ef"],
    "geography": ["geography", "region", "country", "location", "geo", "countrycode"],
    "gwp_set": ["gwpset", "ar", "gwpversion", "ipccversion", "assessmentreport"],
    "source": ["source", "publisher", "database", "datasource", "provider"],
    "source_year": ["sourceyear", "year", "referenceyear", "datayear", "publicationyear"],
    "dataset_ref": ["datasetref", "dataset", "reference", "ref", "uuid", "datasetid", "id"],
    "licence": ["licence", "license", "licencetype"],
    "recycled_variant": ["recycledvariant", "recycled", "secondary", "iscycled", "isrecycled"],
    "valid_from": ["validfrom", "from", "startdate", "validitystart"],
    "valid_to": ["validto", "to", "enddate", "validuntil", "validitiyend", "validityend", "expiry"],
}

# Words in a group cell or a material name → MaterialFactor group.
GROUP_WORDS: list[tuple[str, tuple[str, ...]]] = [
    ("transport", ("truck", "lorry", "ship", "vessel", "freight", "rail", "train", "barge", "air cargo", "airfreight", "shipping", "transport", "hgv", "container ship")),
    ("waste", ("landfill", "incinerat", "waste", "scrap treatment", "disposal", "wastewater")),
    ("energy", ("electricity", "grid", "natural gas", "diesel", "fuel oil", "lpg", "steam", "heat", "energy", "power", "petrol", "gasoline", "kerosene")),
    ("packaging", ("pallet", "drum", "reel", "spool", "carton", "cardboard", "corrugated", "box", "crate", "stretch film", "shrink", "wrap", "packaging", "strap", "bobbin")),
    ("aluminium", ("aluminium", "aluminum", "alu ", "al ", "ec grade", "bauxite", "alumina")),
    ("copper", ("copper", "cu ", "brass", "bronze")),
    ("steel", ("steel", "iron", "galvanised", "galvanized", "stainless")),
    ("polymer", ("pvc", "xlpe", "hdpe", "ldpe", "lldpe", "polyethylene", "polypropylene", "polymer", "plastic", "nylon", "polyamide", "pet ", "rubber", "epdm", "resin", "compound")),
    ("chemical", ("chemical", "lubricant", "oil", "acid", "solvent", "caustic", "lime", "additive", "grease", "coolant", "flux")),
]

LICENCE_ALIASES = {"open": "open", "free": "open", "public": "open", "supplier": "supplier", "epd": "supplier", "ecoinvent": "ecoinvent", "licensed": "ecoinvent", "licenced": "ecoinvent", "proprietary": "ecoinvent"}
TRUE_WORDS = {"yes", "y", "true", "1", "recycled", "secondary", "x"}
FALSE_WORDS = {"no", "n", "false", "0", "primary", "virgin", ""}


class SheetError(ValueError):
    """The file can't be read as a factor sheet; the message is for people."""


def norm(header: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower().replace("₂", "2"))


def clean_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


# ------------------------------------------------------------------ reading ---

@dataclass
class Grid:
    sheet_names: list[str]
    sheet_name: Optional[str]
    rows: list[list[Any]]  # every row, values as read


def read_grid(content: bytes, filename: str, sheet_name: Optional[str] = None) -> Grid:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return Grid([], None, [list(r) for r in csv.reader(io.StringIO(text), dialect)])
    if ext != "xlsx":
        raise SheetError("Use an .xlsx or .csv file. Save an .xls file as .xlsx first.")
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise SheetError("Couldn't open that file as an Excel workbook.") from e
    names = list(book.sheetnames)
    if sheet_name is not None and sheet_name not in names:
        raise SheetError(f"The workbook has no sheet called “{sheet_name}”.")
    grids = {}
    for name in [sheet_name] if sheet_name else names:
        grids[name] = [list(r) for r in book[name].iter_rows(values_only=True)]
    book.close()
    if sheet_name:
        return Grid(names, sheet_name, grids[sheet_name])
    # The first sheet whose best header row names the required columns; else the first with data.
    best = max(names, key=lambda n: (_header_score(grids[n])[1], -names.index(n)))
    return Grid(names, best, grids[best])


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _header_score(rows: list[list[Any]]) -> tuple[int, float]:
    """(0-based header row, score). Score: required fields named, then other fields, then text cells."""
    best_row, best = 0, -1.0
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        texts = [clean_header(c) for c in row if isinstance(c, str) and c.strip()]
        if len(texts) < 2:
            continue
        mapping = match_headers(texts)
        score = sum(10 for f in NEEDS_A_COLUMN if f in mapping) + len(mapping) + len(texts) * 0.01
        if score > best:
            best_row, best = i, score
    return best_row, max(best, 0.0)


# ------------------------------------------------------------------ columns ---

@dataclass
class ColumnMatch:
    header: str
    confidence: int
    reason: str
    by: str  # "header" | "ai" | "person"


def match_headers(headers: list[str]) -> dict[str, ColumnMatch]:
    """Header → field: exact names and synonyms first, then a header containing one, then a close spelling.

    Each pass runs over every field before the next, so “Material group” goes to
    the group (exact) before the name (contains “material”) can take it.
    """
    out: dict[str, ColumnMatch] = {}
    keyed = [(h, norm(h)) for h in headers if norm(h)]

    def free():
        used = {m.header for m in out.values()}
        return [(h, n) for h, n in keyed if h not in used]

    for f, label, _ in FIELDS:
        hit = next((h for h, n in free() if n in SYNONYMS[f]), None)
        if hit:
            out[f] = ColumnMatch(hit, 95, f"Header “{hit}” is the {label.lower()} column", "header")
    for f, label, _ in FIELDS:
        if f in out:
            continue
        hit = next(((h, s) for h, n in free() for s in SYNONYMS[f] if (len(s) >= 5 and s in n) or (len(s) == 4 and n.endswith(s))), None)
        if hit:
            out[f] = ColumnMatch(hit[0], 80, f"Header “{hit[0]}” contains “{hit[1]}”", "header")
        elif f == "value_kgco2e":
            # “GWP (t CO2e / t)”, “CO2-eq”: a header naming CO2e or GWP is the value.
            hit = next((h for h, n in free() if "co2" in n or n.startswith("gwp")), None)
            if hit:
                out[f] = ColumnMatch(hit, 80, f"Header “{hit}” names CO2e", "header")
    for f, label, _ in FIELDS:
        if f in out:
            continue
        scored = [(fuzz.ratio(n, s), h) for h, n in free() for s in SYNONYMS[f] if len(s) >= 5]
        top = max(scored, default=(0, ""))
        if top[0] >= 85:
            out[f] = ColumnMatch(top[1], 70, f"Header “{top[1]}” looks like “{label}”", "header")
    return out


AI_MAPPING_PROMPT = """\
You map spreadsheet columns to the fields of a product-carbon-footprint material factor.
Fields: {fields}.
The sheet's headers and a few rows follow as JSON. Return ONLY a JSON object:
{{"mapping": {{"<field>": {{"header": "<exact header>", "confidence": <0-100>, "reason": "<one short sentence>"}}}}}}
Use only headers that exist, at most once each. Leave out a field you can't place.
value_kgco2e is the numeric emission factor (kg, g or t of CO2e per unit of the material); never a year, a quantity or a price.
"""


def ai_mapping(
    headers: list[str],
    sample: list[dict[str, Any]],
    missing: list[str],
    call_llm: Callable[[list[dict]], str],
) -> dict[str, ColumnMatch]:
    from app.services.llm import _parse_json_object

    fields = ", ".join(f"{f} ({FIELD_LABEL[f]})" for f in missing)
    messages = [
        {"role": "system", "content": AI_MAPPING_PROMPT.format(fields=fields)},
        {"role": "user", "content": json.dumps({"headers": headers, "rows": sample}, default=str)[:12000]},
    ]
    data = _parse_json_object(call_llm(messages))
    raw = data.get("mapping") if isinstance(data, dict) else None
    out: dict[str, ColumnMatch] = {}
    if not isinstance(raw, dict):
        return out
    for f in missing:
        m = raw.get(f)
        if not isinstance(m, dict) or m.get("header") not in headers:
            continue
        try:
            conf = int(max(0, min(90, float(m.get("confidence", 50)))))  # AI never outranks an exact header
        except (TypeError, ValueError):
            conf = 50
        reason = str(m.get("reason") or "Picked by the AI sheet reader").strip()[:200]
        out[f] = ColumnMatch(m["header"], conf, reason, "ai")
    return out


# --------------------------------------------------------------------- rows ---

def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def _number(v: Any) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = _cell_text(v).replace(" ", "")
    if not t:
        return None
    # "0,5" is a decimal comma; "1,000" and "1,234,567" are thousands.
    if re.fullmatch(r"-?\d+,\d+", t) and not re.fullmatch(r"-?\d{1,3}(,\d{3})+", t):
        t = t.replace(",", ".")
    t = t.replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


def _date(v: Any) -> Optional[str]:
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    t = _cell_text(v)
    if not t:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(t[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return t  # ESG-lite refuses it with its own message


def infer_group(*texts: str) -> Optional[tuple[str, str]]:
    """(group, the word that decided it) from a group cell or a material name."""
    for text in texts:
        if not text:
            continue
        low = f" {text.lower()} "
        if text.strip().lower() in GROUPS:
            return text.strip().lower(), text.strip()
        for group, words in GROUP_WORDS:
            for w in words:
                if re.search(r"(?<![a-z])" + re.escape(w), low):
                    return group, w.strip()
    return None


@dataclass
class ParsedRow:
    line: int
    values: dict[str, Any]
    confidence: int
    reason: str
    warnings: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)


def parse_row(
    line: int,
    raw: dict[str, Any],
    mapping: dict[str, ColumnMatch],
    value_scale: Optional[tuple[float, str]],
    sheet_unit: Optional[PcfUnit],
) -> ParsedRow:
    get = lambda f: raw.get(mapping[f].header) if f in mapping else None  # noqa: E731
    v: dict[str, Any] = {}
    warnings: list[str] = []
    problems: list[str] = []
    confs = [mapping[f].confidence for f in NEEDS_A_COLUMN if f in mapping]
    notes: list[str] = []

    name = _cell_text(get("name"))
    v["name"] = name
    if not name:
        problems.append("Name is empty")

    # Group: the cell, read loosely, else the name.
    group_cell = _cell_text(get("material_group"))
    g = infer_group(group_cell) if group_cell else None
    if g:
        v["material_group"] = g[0]
        if group_cell.lower() != g[0]:
            notes.append(f"group “{group_cell}” read as {g[0]}")
        confs.append(min(mapping["material_group"].confidence, 80 if group_cell.lower() != g[0] else 100))
    else:
        g = infer_group(name)
        v["material_group"] = g[0] if g else ""
        if g:
            notes.append(f"group {g[0]} taken from “{g[1]}” in the name")
            confs.append(65)
            if group_cell:
                warnings.append(f"Group “{group_cell}” is not a material group; used {g[0]} from the name")
        elif group_cell:
            problems.append(f"Group “{group_cell}” is not one of {', '.join(GROUPS)}")
        else:
            problems.append(f"Group is empty and can't be told from the name; use one of {', '.join(GROUPS)}")

    # Unit: the cell (or the value header's “per …”), converted to a PCF unit.
    unit_cell = _cell_text(get("unit"))
    unit = to_pcf_unit(unit_cell) if unit_cell else sheet_unit
    multiplier = 1.0
    if unit is None:
        v["unit"] = ""
        problems.append("Unit is empty")
    elif isinstance(unit, PcfUnit):
        v["unit"] = unit.unit
        multiplier = 1 / unit.per_written
        if unit.per_written != 1.0 or (unit_cell and unit_cell != unit.unit):
            notes.append(unit.reason)
        confs.append(mapping["unit"].confidence if unit_cell else 75)
    else:
        v["unit"] = unit_cell
        problems.append(unit.reason)

    # Value: kgCO2e per unit, scaled from tCO2e/gCO2e and to the PCF unit.
    cell_scale = co2e_multiplier(unit_cell) if unit_cell else None
    scale = cell_scale or value_scale
    number = _number(get("value_kgco2e"))
    if number is None or number < 0:
        problems.append("Value must be a number of 0 or more")
        v["value_kgco2e"] = None
    else:
        factor = (scale[0] if scale else 1.0) * multiplier
        v["value_kgco2e"] = round(number * factor, 12)
        if scale and scale[0] != 1.0:
            notes.append(scale[1])

    for f in ("geography", "source", "dataset_ref"):
        if f in mapping:
            v[f] = _cell_text(get(f)) or None
    if "source_year" in mapping:
        t = _cell_text(get("source_year"))
        n = _number(t)
        if not t:
            v["source_year"] = None
        elif n is not None and n.is_integer() and 1900 <= n <= 2100:
            v["source_year"] = int(n)
        else:
            m = re.search(r"\b(19|20)\d{2}\b", t)
            if m:
                v["source_year"] = int(m.group(0))
                notes.append(f"year {m.group(0)} taken from “{t}”")
            else:
                problems.append("Source year must be a year")
    if "gwp_set" in mapping:
        t = _cell_text(get("gwp_set")).upper()
        m = re.search(r"AR\s*([56])", t)
        if m:
            v["gwp_set"] = f"AR{m.group(1)}"
        elif t:
            problems.append("GWP set must be AR6 or AR5")
    if "licence" in mapping:
        t = _cell_text(get("licence")).lower()
        lic = next((val for k, val in LICENCE_ALIASES.items() if k in t), None) if t else None
        if lic:
            v["licence"] = lic
        elif t:
            problems.append("Licence must be open, supplier or ecoinvent")
    if "recycled_variant" in mapping:
        t = _cell_text(get("recycled_variant")).lower()
        if t in TRUE_WORDS:
            v["recycled_variant"] = True
        elif t in FALSE_WORDS:
            v["recycled_variant"] = False
        else:
            v["recycled_variant"] = False
            warnings.append(f"Recycled “{t}” read as no")
    for f in ("valid_from", "valid_to"):
        if f in mapping:
            v[f] = _date(get(f))

    by_ai = [FIELD_LABEL[f].split(" (")[0] for f, m in mapping.items() if m.by == "ai"]
    lead = f"AI matched {', '.join(by_ai)}" if by_ai else "Columns matched by header"
    reason = "; ".join([lead] + notes)
    confidence = min(confs) if confs else 0
    return ParsedRow(line, v, int(confidence), reason, warnings, problems)


# ------------------------------------------------------------------- driver ---

def _json_safe(v: Any) -> Any:
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def read_factor_sheet(
    content: bytes,
    filename: str,
    sheet_name: Optional[str] = None,
    mapping_override: Optional[dict[str, Optional[str]]] = None,
    call_llm: Optional[Callable[[list[dict]], str]] = None,
) -> dict[str, Any]:
    grid = read_grid(content, filename, sheet_name)
    header_idx, _ = _header_score(grid.rows)
    if not grid.rows or header_idx >= len(grid.rows):
        raise SheetError("The sheet is empty.")
    header_cells = grid.rows[header_idx]
    columns = [(i, clean_header(c)) for i, c in enumerate(header_cells) if not _blank(c)]
    headers = [h for _, h in columns]
    if len(headers) < 2:
        raise SheetError("Couldn't find a header row with at least two columns.")

    data: list[tuple[int, dict[str, Any]]] = []
    for offset, row in enumerate(grid.rows[header_idx + 1:], start=header_idx + 2):
        cells = {h: _json_safe(row[i]) if i < len(row) else None for i, h in columns}
        if all(_blank(v) for v in cells.values()):
            continue
        data.append((offset, cells))
    if not data:
        raise SheetError("The sheet has a header row but no data rows.")
    if len(data) > MAX_ROWS:
        raise SheetError(f"The sheet has {len(data)} rows; read at most {MAX_ROWS} at a time.")

    warnings: list[str] = []
    mapping = match_headers(headers)
    if mapping_override is not None:
        for f, h in mapping_override.items():
            if f not in FIELD_LABEL:
                continue
            if h is None or h == "":
                mapping.pop(f, None)
            elif h in headers:
                # The header moves to this field; drop it from any other.
                for other in [k for k, m in mapping.items() if m.header == h and k != f]:
                    mapping.pop(other)
                mapping[f] = ColumnMatch(h, 100, f"You picked “{h}”", "person")
            else:
                warnings.append(f"The sheet has no column “{h}”; {FIELD_LABEL[f]} was left unmatched")

    ai_used = False
    missing = [f for f in NEEDS_A_COLUMN if f not in mapping]
    if missing and call_llm is not None:
        free = [h for h in headers if h not in {m.header for m in mapping.values()}]
        try:
            found = ai_mapping(free, [r for _, r in data[:SAMPLE_ROWS]], missing, call_llm)
            mapping.update(found)
            ai_used = bool(found)
        except Exception:
            logger.warning("AI column mapping failed; using header names only", exc_info=True)
            warnings.append("The AI sheet reader is unavailable right now; columns were matched by header name only.")

    # A value column that is mostly text is probably the wrong column.
    if "value_kgco2e" in mapping:
        h = mapping["value_kgco2e"].header
        filled = [r[h] for _, r in data if not _blank(r[h])]
        numeric = sum(1 for x in filled if _number(x) is not None)
        if filled and numeric / len(filled) < 0.5:
            m = mapping["value_kgco2e"]
            mapping["value_kgco2e"] = ColumnMatch(m.header, min(m.confidence, 40), f"{m.reason}, but most of its cells are not numbers", m.by)
            warnings.append(f"Most cells in “{h}” are not numbers; check the value column.")

    value_header = mapping["value_kgco2e"].header if "value_kgco2e" in mapping else ""
    value_scale = co2e_multiplier(value_header) if value_header else None
    sheet_unit = None
    if "unit" not in mapping and value_header:
        u = to_pcf_unit(value_header)
        if isinstance(u, PcfUnit):
            sheet_unit = u

    rows = [parse_row(line, raw, mapping, value_scale, sheet_unit) for line, raw in data]
    for f in NEEDS_A_COLUMN:
        if f not in mapping:
            warnings.append(f"No column for {FIELD_LABEL[f]}; pick one.")

    return {
        "file_name": filename,
        "sheet_names": grid.sheet_names,
        "sheet_name": grid.sheet_name,
        "header_line": header_idx + 1,
        "headers": headers,
        "mapping": {f: m.__dict__ for f, m in mapping.items()},
        "sheet_unit": (
            {"unit": sheet_unit.unit, "reason": f"No unit column; every row is per {sheet_unit.unit}, from the value header “{value_header}”"}
            if sheet_unit
            else None
        ),
        "rows": [r.__dict__ for r in rows],
        "total_rows": len(rows),
        "rows_with_problems": sum(1 for r in rows if r.problems),
        "low_confidence_rows": sum(1 for r in rows if r.confidence < 60),
        "warnings": warnings,
        "ai_used": ai_used,
    }
